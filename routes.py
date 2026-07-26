# Изменение №4 (бэкенд-рефакторинг): выделены маршруты в routes.py

import sqlite3
import os
import time
import io
import secrets
import logging
import threading
from datetime import datetime, timedelta
from functools import wraps

logger = logging.getLogger(__name__)
from flask import Blueprint, jsonify, request, session, send_file, render_template
from werkzeug.security import check_password_hash

from config import (
    BELL_SCHEDULE, SHIFT_LABELS, ADMIN_PASSWORD_HASH,
    LOGIN_RATE_LIMIT as _CONFIG_LOGIN_RATE_LIMIT,
    LOGIN_RATE_WINDOW as _CONFIG_LOGIN_RATE_WINDOW,
    BACKUP_RETENTION_DAYS as _CONFIG_BACKUP_RETENTION_DAYS,
    TRASH_RETENTION_DAYS,
    LOGIN_CLEANUP_INTERVAL
)
from db import get_db, is_date_excluded
from utils import _safe_filename, _validate_lesson_data, LESSON_SELECT, LESSON_SELECT_WITH_DAY
from cache import refs_cache

bp = Blueprint('main', __name__)


@bp.after_request
def set_charset(response):
    if response.content_type and 'text/html' in response.content_type:
        response.content_type = 'text/html; charset=utf-8'
    elif response.content_type and 'application/json' in response.content_type:
        response.content_type = 'application/json; charset=utf-8'
    if request.path.startswith('/static/'):
        response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
    return response


_login_attempts = {}
_last_cleanup_time = 0
_login_lock = threading.Lock()

# Мутабельные настройки (инициализируются из config, перезаписываются из БД при старте и через API)
_settings = {
    'session_lifetime': 3600,
    'login_rate_limit': _CONFIG_LOGIN_RATE_LIMIT,
    'login_rate_window': _CONFIG_LOGIN_RATE_WINDOW,
    'backup_retention_days': _CONFIG_BACKUP_RETENTION_DAYS,
}

def _refresh_settings_from_db():
    try:
        conn = get_db()
        rows = conn.execute('SELECT key, value FROM settings').fetchall()
        conn.close()
        for r in rows:
            if r['key'] in _settings:
                _settings[r['key']] = int(r['value'])
    except Exception:
        pass

_refresh_settings_from_db()


def _clean_login_attempts():
    global _last_cleanup_time
    now = time.time()
    if now - _last_cleanup_time < LOGIN_CLEANUP_INTERVAL:
        return
    _last_cleanup_time = now
    expired = [ip for ip, ts in _login_attempts.items() if not ts or now - ts[-1] >= _settings['login_rate_window']]
    for ip in expired:
        _login_attempts.pop(ip, None)


def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get('is_admin'):
            return jsonify({"error": "Unauthorized"}), 401
        return f(*args, **kwargs)
    return decorated


@bp.before_request
def csrf_check():
    if request.method in ('POST', 'PUT', 'DELETE', 'PATCH'):
        if request.path.startswith('/api/'):
            if request.path in ('/api/auth/login', '/api/auth/logout', '/api/auth/check'):
                return
            token = request.headers.get('X-CSRF-Token') or request.headers.get('X-CSRF-TOKEN') or ''
            if token != session.get('csrf_token'):
                return jsonify({"error": "CSRF token missing or invalid"}), 403


def build_week_dates(monday_str):
    monday = datetime.strptime(monday_str, '%Y-%m-%d').date()
    names = ['Понедельник', 'Вторник', 'Среда', 'Четверг', 'Пятница', 'Суббота']
    return [{'day_name': names[i], 'date': (monday + timedelta(days=i)).isoformat()} for i in range(6)]


def _fetch_group_raw(group_name, monday):
    week_dates = build_week_dates(monday)
    date_list = [d['date'] for d in week_dates]
    conn = get_db()
    placeholders = ','.join('?' * len(date_list))
    days = conn.execute(
        f'SELECT * FROM schedule_days WHERE group_name = ? AND deleted_at IS NULL AND date IN ({placeholders}) ORDER BY date',
        [group_name] + date_list
    ).fetchall()
    day_id_map = {day['date']: day['id'] for day in days}
    for wd in week_dates:
        wd['day_id'] = day_id_map.get(wd['date'])
    lessons_by_date = {}
    for day in days:
        lessons = conn.execute(
            LESSON_SELECT + ' WHERE l.day_id = ? AND l.deleted_at IS NULL ORDER BY l.period',
            (day['id'],)
        ).fetchall()
        lessons_by_date[day['date']] = {}
        for l in lessons:
            p = l['period']
            if p not in lessons_by_date[day['date']]:
                lessons_by_date[day['date']][p] = []
            lessons_by_date[day['date']][p].append(dict(l))
    conn.close()
    return week_dates, lessons_by_date


def _fetch_entity_raw(entity_table, entity_id, entity_col, monday):
    week_dates = build_week_dates(monday)
    date_list = [d['date'] for d in week_dates]
    conn = get_db()
    placeholders = ','.join('?' * len(date_list))
    sql = f'''
        {LESSON_SELECT_WITH_DAY}
        WHERE l.{entity_col} = ? AND l.subject_id IS NOT NULL AND l.deleted_at IS NULL AND sd.deleted_at IS NULL AND sd.date IN ({placeholders})
        ORDER BY sd.date, l.period
    '''
    params = [entity_id] + date_list
    logger.debug("Export SQL: %s | params: %s", sql.strip().replace('\n', ' '), params)
    rows = conn.execute(sql, params).fetchall()
    logger.debug("Export rows count: %d", len(rows))
    conn.close()
    lessons_by_date_period = {}
    day_id_map = {}
    for r in rows:
        d = r['date']
        p = r['period']
        day_id_map[d] = r['day_id']
        if d not in lessons_by_date_period:
            lessons_by_date_period[d] = {}
        lessons_by_date_period[d][p] = {
            'id': r['id'], 'subject': r['subject_name'], 'teacher': r['teacher_name'],
            'classroom': r['classroom_name'],
            'groups': r['group_name'] or r['groups'],
            'day_id': r['day_id'],
            'group_id': r['group_id'], 'subgroup': r['subgroup']
        }
    for wd in week_dates:
        wd['day_id'] = day_id_map.get(wd['date'])
    return week_dates, lessons_by_date_period


def _build_schedule_periods(week_dates, lessons_by_date_period):
    periods = []
    for b in BELL_SCHEDULE:
        p = b['period']
        lessons = {}
        for wd in week_dates:
            d = wd['date']
            lesson = lessons_by_date_period.get(d, {}).get(p)
            if lesson:
                lessons[d] = [{
                    'id': lesson['id'], 'subject': lesson['subject'], 'teacher': lesson['teacher'],
                    'classroom': lesson['classroom'], 'groups': lesson['groups'],
                    'day_id': lesson['day_id'],
                    'group_id': lesson.get('group_id'), 'subgroup': lesson.get('subgroup')
                }]
            else:
                lessons[d] = []
        periods.append({**b, 'lessons': lessons})
    return periods


def _get_export_data(schedule_type, name, monday, entity_id=None):
    if schedule_type == 'group':
        week_dates, lessons_by_date = _fetch_group_raw(name, monday)
        periods = []
        for b in BELL_SCHEDULE:
            p = b['period']
            lessons = {}
            for wd in week_dates:
                d = wd['date']
                if d in lessons_by_date and p in lessons_by_date[d]:
                    raw_list = lessons_by_date[d][p]
                    lessons[d] = '\n---\n'.join(
                        f"{l['subject_name'] or ''}\n{l['teacher_name'] or ''}\n{l['classroom_name'] or ''}"
                        for l in raw_list
                    )
                else:
                    lessons[d] = ''
            periods.append({**b, 'lessons': lessons})
        return week_dates, periods

    elif schedule_type == 'teacher':
        tid = entity_id
        if not tid:
            logger.warning("Export teacher: no id provided, fallback to name lookup for '%s'", name)
            conn = get_db()
            row = conn.execute(
                "SELECT id FROM teachers WHERE REPLACE(name, ' ', '') LIKE REPLACE(?, ' ', '') || '%'",
                (name,)
            ).fetchone()
            conn.close()
            if not row:
                return build_week_dates(monday), []
            tid = row['id']
        logger.info("Export teacher: id=%d", tid)
        week_dates, ldp = _fetch_entity_raw('teachers', tid, 'teacher_id', monday)
        periods = []
        for b in BELL_SCHEDULE:
            p = b['period']
            lessons = {}
            for wd in week_dates:
                l = ldp.get(wd['date'], {}).get(p)
                lessons[wd['date']] = f"{l['subject'] or ''}\n{l['groups'] or ''}\n{l['classroom'] or ''}" if l else ''
            periods.append({**b, 'lessons': lessons})
        return week_dates, periods

    elif schedule_type == 'classroom':
        cid = entity_id
        if not cid:
            logger.warning("Export classroom: no id provided, fallback to name lookup for '%s'", name)
            conn = get_db()
            row = conn.execute(
                "SELECT id FROM classrooms_list WHERE REPLACE(name, ' ', '') LIKE REPLACE(?, ' ', '') || '%'",
                (name,)
            ).fetchone()
            conn.close()
            if not row:
                return build_week_dates(monday), []
            cid = row['id']
        logger.info("Export classroom: id=%d", cid)
        week_dates, ldp = _fetch_entity_raw('classrooms_list', cid, 'classroom_id', monday)
        periods = []
        for b in BELL_SCHEDULE:
            p = b['period']
            lessons = {}
            for wd in week_dates:
                l = ldp.get(wd['date'], {}).get(p)
                lessons[wd['date']] = f"{l['subject'] or ''}\n{l['teacher'] or ''}\n{l['groups'] or ''}" if l else ''
            periods.append({**b, 'lessons': lessons})
        return week_dates, periods

    return build_week_dates(monday), []


def _export_excel(week_dates, periods, title):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = 'Расписание'

    thin = Side(style='thin', color='999999')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(name='Arial', bold=True, size=11)
    header_fill = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')
    cell_font = Font(name='Arial', size=10)
    wrap = Alignment(wrap_text=True, vertical='top')

    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(week_dates) + 1)
    title_cell = ws.cell(row=row, column=1, value=title)
    title_cell.font = Font(name='Arial', bold=True, size=14)
    row += 1

    ws.cell(row=row, column=1, value='Пара').font = header_font
    ws.cell(row=row, column=1).fill = header_fill
    ws.cell(row=row, column=1).border = border
    ws.cell(row=row, column=1).alignment = Alignment(horizontal='center')

    for i, wd in enumerate(week_dates):
        d = datetime.strptime(wd['date'], '%Y-%m-%d')
        header_text = f"{wd['day_name']}\n{d.day}.{d.month}"
        cell = ws.cell(row=row, column=i + 2, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
    row += 1

    last_shift = 0
    for p in periods:
        if p['shift'] != last_shift:
            last_shift = p['shift']
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(week_dates) + 1)
            shift_cell = ws.cell(row=row, column=1, value=SHIFT_LABELS[p['shift']])
            shift_cell.font = Font(name='Arial', bold=True, size=10, color='1A5FB4')
            shift_cell.fill = PatternFill(start_color='E0ECFF', end_color='E0ECFF', fill_type='solid')
            shift_cell.alignment = Alignment(horizontal='center')
            for c in range(1, len(week_dates) + 2):
                ws.cell(row=row, column=c).border = border
            row += 1

        time_cell = ws.cell(row=row, column=1, value=f"{p['period']} пара\n{p['start']}–{p['end']}")
        time_cell.font = cell_font
        time_cell.border = border
        time_cell.alignment = Alignment(horizontal='center', vertical='top')

        for i, wd in enumerate(week_dates):
            cell = ws.cell(row=row, column=i + 2, value=p['lessons'].get(wd['date'], ''))
            cell.font = cell_font
            cell.border = border
            cell.alignment = wrap
        row += 1

    ws.column_dimensions['A'].width = 16
    for i in range(len(week_dates)):
        letter = get_column_letter(i + 2)
        ws.column_dimensions[letter].width = 28

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _export_word(week_dates, periods, title):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii', 'Arial')
    style.element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi', 'Arial')

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(16)

    num_cols = len(week_dates) + 1
    table = doc.add_table(rows=1, cols=num_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0]
    hdr.cells[0].text = 'Пара'
    for p in hdr.cells[0].paragraphs:
        for r in p.runs:
            r.bold = True
    for i, wd in enumerate(week_dates):
        d = datetime.strptime(wd['date'], '%Y-%m-%d')
        cell = hdr.cells[i + 1]
        cell.text = f"{wd['day_name']}\n{d.day}.{d.month}"
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True

    last_shift = 0
    for p in periods:
        if p['shift'] != last_shift:
            last_shift = p['shift']
            shift_row = table.add_row()
            shift_row.cells[0].merge(shift_row.cells[len(week_dates)])
            shift_row.cells[0].text = SHIFT_LABELS[p['shift']]
            for para in shift_row.cells[0].paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x1A, 0x5F, 0xB4)

        row_cells = table.add_row().cells
        row_cells[0].text = f"{p['period']} пара\n{p['start']}–{p['end']}"
        for i, wd in enumerate(week_dates):
            row_cells[i + 1].text = p['lessons'].get(wd['date'], '')

    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs:
                    run.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _export_pdf(week_dates, periods, title):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_path = os.environ.get('PDF_FONT_PATH', '')
    if not font_path:
        candidates = [
            r'C:\Windows\Fonts\tahoma.ttf',
            r'C:\Windows\Fonts\arial.ttf',
            r'C:\Windows\Fonts\segoeui.ttf',
            '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
            '/usr/share/fonts/TTF/DejaVuSans.ttf',
        ]
        for cp in candidates:
            if os.path.exists(cp):
                font_path = cp
                break
    if not font_path or not os.path.exists(font_path):
        font_path = None

    font_bold_path = os.environ.get('PDF_FONT_BOLD_PATH', '')
    if not font_bold_path and font_path:
        bname = font_path.replace('.ttf', 'bd.ttf')
        if os.path.exists(bname):
            font_bold_path = bname

    if font_path:
        pdfmetrics.registerFont(TTFont('Tahoma', font_path))
        pdfmetrics.registerFont(TTFont('Tahoma-Bold', font_bold_path or font_path))
    font_name = 'Tahoma' if font_path else 'Helvetica'
    font_bold_name = 'Tahoma-Bold' if font_path else 'Helvetica-Bold'

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)

    style_title = ParagraphStyle('Title', fontName=font_bold_name, fontSize=14, alignment=1, spaceAfter=12)
    style_cell = ParagraphStyle('Cell', fontName=font_name, fontSize=8, leading=10)
    style_header = ParagraphStyle('Header', fontName=font_bold_name, fontSize=9, leading=11, alignment=1)

    elements = []
    elements.append(Paragraph(title, style_title))
    elements.append(Spacer(1, 0.3*cm))

    header_dates = [(wd['day_name'], datetime.strptime(wd['date'], '%Y-%m-%d')) for wd in week_dates]
    data = [['Пара'] + [f"{name}\n{d.day}.{d.month}" for name, d in header_dates]]

    last_shift = 0
    for p in periods:
        if p['shift'] != last_shift:
            last_shift = p['shift']
            shift_row_data = [Paragraph(f"<b>{SHIFT_LABELS[p['shift']]}</b>", style_header)]
            shift_row_data += [''] * len(week_dates)
            data.append(shift_row_data)
        row = [Paragraph(f"<b>{p['period']} пара</b><br/>{p['start']}–{p['end']}", style_cell)]
        for wd in week_dates:
            text = p['lessons'].get(wd['date'], '')
            row.append(Paragraph(text.replace('\n', '<br/>'), style_cell) if text else '')
        data.append(row)

    col_widths = [2.5*cm] + [max(2.2*cm, (landscape(A4)[0] - 4*cm - 2.5*cm) / len(week_dates))] * len(week_dates)
    t = Table(data, colWidths=col_widths, repeatRows=1)

    style_cmds = [
        ('FONTNAME', (0, 0), (-1, 0), font_bold_name),
        ('FONTNAME', (0, 1), (-1, -1), font_name),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#D9E2F3')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#999999')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]

    row_idx = 1
    last_shift = 0
    for p in periods:
        if p['shift'] != last_shift:
            last_shift = p['shift']
            style_cmds.append(('SPAN', (0, row_idx), (-1, row_idx)))
            style_cmds.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor('#E0ECFF')))
            style_cmds.append(('ALIGN', (0, row_idx), (-1, row_idx), 'CENTER'))
            row_idx += 1
        row_idx += 1

    t.setStyle(TableStyle(style_cmds))
    elements.append(t)
    doc.build(elements)
    buf.seek(0)
    return buf


def generic_crud(table, label, order_by='id', extra_fields=None):
    extra_fields = extra_fields or []

    def make_get(label, table, order_by):
        def get_items():
            """Список записей справочника
            ---
            responses:
              200:
                description: Массив объектов справочника
            """
            cache_key = f'ref_list_{table}'
            cached = refs_cache.get(cache_key)
            if cached is not None:
                return jsonify(cached)
            conn = get_db()
            items = conn.execute(f"SELECT * FROM {table} WHERE deleted_at IS NULL ORDER BY {order_by}").fetchall()
            conn.close()
            result = [dict(i) for i in items]
            refs_cache.set(cache_key, result)
            return jsonify(result)
        get_items.__name__ = f'get_{label}'
        return get_items

    def make_create(label, table, extra_fields):
        @login_required
        def create_item():
            """Создать запись в справочнике
            ---
            parameters:
              - in: body
                name: body
                required: true
                schema:
                  type: object
                  properties:
                    name:
                      type: string
                      description: Название
            responses:
              201:
                description: Созданная запись
              400:
                description: Name required
              409:
                description: Уже существует
            """
            data = request.get_json()
            name = data.get('name', '').strip()
            if not name:
                return jsonify({"error": "Name required"}), 400
            conn = get_db()
            extra_cols = []
            extra_vals = []
            for f in extra_fields:
                if f in data:
                    extra_cols.append(f)
                    extra_vals.append(data[f])
            if table == 'teachers' and 'phone' in data:
                phone = (data['phone'] or '').strip()
                if phone and conn.execute("SELECT 1 FROM teachers WHERE phone = ? AND deleted_at IS NULL", (phone,)).fetchone():
                    conn.close()
                    return jsonify({"error": "Этот номер телефона уже используется другим преподавателем"}), 409
            cols_sql = ', '.join(['name'] + extra_cols)
            placeholders = ', '.join(['?'] * (1 + len(extra_cols)))
            try:
                cursor = conn.execute(f'INSERT INTO {table} ({cols_sql}) VALUES ({placeholders})', [name] + extra_vals)
                conn.commit()
                item_id = cursor.lastrowid
            except sqlite3.IntegrityError:
                conn.close()
                return jsonify({"error": "Уже существует"}), 409
            conn.close()
            result = {"id": item_id, "name": name}
            for f in extra_fields:
                result[f] = data.get(f)
            refs_cache.invalidate_prefix(f'ref_list_{table}')
            return jsonify(result), 201
        create_item.__name__ = f'create_{label}'
        return create_item

    def make_update(label, table, extra_fields):
        @login_required
        def update_item(item_id):
            """Обновить запись в справочнике
            ---
            parameters:
              - in: path
                name: item_id
                type: integer
                required: true
              - in: body
                name: body
                required: true
                schema:
                  type: object
                  properties:
                    name:
                      type: string
            responses:
              200:
                description: success
              400:
                description: Name required
            """
            data = request.get_json()
            name = data.get('name', '').strip()
            if not name:
                return jsonify({"error": "Name required"}), 400
            conn = get_db()
            if table == 'teachers' and 'phone' in data:
                phone = (data['phone'] or '').strip()
                if phone and conn.execute("SELECT 1 FROM teachers WHERE phone = ? AND id != ? AND deleted_at IS NULL", (phone, item_id)).fetchone():
                    conn.close()
                    return jsonify({"error": "Этот номер телефона уже используется другим преподавателем"}), 409
            set_parts = ['name = ?']
            params = [name]
            for f in extra_fields:
                if f in data:
                    set_parts.append(f'{f} = ?')
                    params.append(data[f])
            params.append(item_id)
            conn.execute(f'UPDATE {table} SET {", ".join(set_parts)} WHERE id = ?', params)
            conn.commit()
            conn.close()
            refs_cache.invalidate_prefix(f'ref_list_{table}')
            return jsonify({"success": True})
        update_item.__name__ = f'update_{label}'
        return update_item

    def make_delete(label, table):
        @login_required
        def delete_item(item_id):
            """Мягкое удаление записи из справочника
            ---
            parameters:
              - in: path
                name: item_id
                type: integer
                required: true
            responses:
              200:
                description: success (установлен deleted_at)
            """
            conn = get_db()
            conn.execute(f"UPDATE {table} SET deleted_at = datetime('now') WHERE id = ? AND deleted_at IS NULL", (item_id,))
            conn.commit()
            conn.close()
            refs_cache.invalidate_prefix(f'ref_list_{table}')
            return jsonify({"success": True, "warning": "Запись помечена как удалённая. Она может использоваться в расписании."})
        delete_item.__name__ = f'delete_{label}'
        return delete_item

    bp.add_url_rule(f'/api/{label}', f'get_{label}', make_get(label, table, order_by), methods=['GET'])
    bp.add_url_rule(f'/api/{label}', f'create_{label}', make_create(label, table, extra_fields), methods=['POST'])
    bp.add_url_rule(f'/api/{label}/<int:item_id>', f'update_{label}', make_update(label, table, extra_fields), methods=['PUT'])
    bp.add_url_rule(f'/api/{label}/<int:item_id>', f'delete_{label}', make_delete(label, table), methods=['DELETE'])


@bp.route('/')
def index():
    """Главная страница — SPA
    ---
    responses:
      200:
        description: HTML-шаблон страницы
    """
    return render_template('index.html')


@bp.route('/api/bell-schedule')
def bell_schedule():
    """Расписание звонков (8 пар, 3 смены)
    ---
    responses:
      200:
        description: Массив периодов с полями period, start, end, shift
    """
    cached = refs_cache.get('bell_schedule')
    if cached is not None:
        return jsonify(cached)
    refs_cache.set('bell_schedule', BELL_SCHEDULE)
    return jsonify(BELL_SCHEDULE)


@bp.route('/api/auth/csrf-token')
def csrf_token():
    """Получить CSRF-токен для защиты форм
    ---
    responses:
      200:
        description: Объект с csrf_token
    """
    if 'csrf_token' not in session:
        session['csrf_token'] = secrets.token_hex(32)
    return jsonify({"csrf_token": session['csrf_token']})


@bp.route('/api/auth/login', methods=['POST'])
def login():
    """Вход администратора по паролю
    ---
    parameters:
      - in: body
        name: body
        required: true
        schema:
          type: object
          properties:
            password:
              type: string
              description: Пароль из .env
    responses:
      200:
        description: Успешный вход, возвращает csrf_token
      401:
        description: Неверный пароль
      429:
        description: Слишком много попыток
    """
    _clean_login_attempts()
    data = request.get_json()
    ip = request.remote_addr
    now = time.time()

    with _login_lock:
        attempts = _login_attempts.get(ip, [])
        attempts = [t for t in attempts if now - t < _settings['login_rate_window']]
        _login_attempts[ip] = attempts

        if len(attempts) >= _settings['login_rate_limit']:
            return jsonify({"error": "Too many attempts. Try later."}), 429

        if data.get('password') and check_password_hash(ADMIN_PASSWORD_HASH, data['password']):
            session['is_admin'] = True
            session['csrf_token'] = secrets.token_hex(32)
            _login_attempts.pop(ip, None)
            return jsonify({"success": True, "csrf_token": session['csrf_token']})

        attempts.append(now)
        _login_attempts[ip] = attempts
    return jsonify({"error": "Wrong password"}), 401


@bp.route('/api/auth/logout', methods=['POST'])
def logout():
    """Выход администратора
    ---
    responses:
      200:
        description: Успешный выход
    """
    session.pop('is_admin', None)
    return jsonify({"success": True})


@bp.route('/api/auth/check')
def check_auth():
    """Проверить статус администратора
    ---
    responses:
      200:
        description: Объект с полем is_admin (bool)
    """
    return jsonify({"is_admin": session.get('is_admin', False)})


@bp.route('/api/settings', methods=['GET'])
@login_required
def get_settings():
    """Получить все настройки
    ---
    responses:
      200:
        description: Объект ключ-значение
    """
    conn = get_db()
    rows = conn.execute('SELECT key, value FROM settings ORDER BY key').fetchall()
    conn.close()
    return jsonify({r['key']: r['value'] for r in rows})


@bp.route('/api/settings/<key>', methods=['PUT'])
@login_required
def update_setting(key):
    """Обновить настройку
    ---
    parameters:
      - in: path
        name: key
        required: true
        type: string
      - in: body
        name: body
        required: true
        schema:
          type: object
          properties:
            value:
              type: string
    responses:
      200:
        description: success
      400:
        description: Неверное значение
      403:
        description: Неверный ключ
    """
    allowed = ('session_lifetime', 'login_rate_limit', 'login_rate_window', 'backup_retention_days')
    if key not in allowed:
        return jsonify({"error": "Unknown setting"}), 403
    data = request.get_json()
    val = data.get('value', '').strip()
    if not val.isdigit() or int(val) < 1:
        return jsonify({"error": "Значение должно быть положительным целым числом"}), 400
    conn = get_db()
    conn.execute("INSERT INTO settings (key, value, updated_at) VALUES (?, ?, CURRENT_TIMESTAMP) ON CONFLICT(key) DO UPDATE SET value = excluded.value, updated_at = CURRENT_TIMESTAMP", (key, val))
    conn.commit()
    conn.close()
    _settings[key] = int(val)
    if key == 'session_lifetime':
        from flask import current_app
        current_app.permanent_session_lifetime = timedelta(seconds=int(val))
    return jsonify({"success": True})


@bp.route('/api/check_date')
def check_date():
    """Проверить, является ли дата исключённой (праздник/выходной)
    ---
    parameters:
      - in: query
        name: date
        type: string
        required: true
        description: Дата в формате YYYY-MM-DD
    responses:
      200:
        description: Объект с полями date и excluded (bool)
      400:
        description: date обязателен
    """
    date_str = request.args.get('date', '')
    if not date_str:
        return jsonify({"error": "date parameter required (YYYY-MM-DD)"}), 400
    conn = get_db()
    excluded = is_date_excluded(conn, date_str)
    conn.close()
    return jsonify({"date": date_str, "excluded": excluded})


@bp.route('/api/days', methods=['GET'])
def get_days():
    """Список дней расписания
    ---
    parameters:
      - in: query
        name: group
        type: string
        required: false
        description: Фильтр по группе
    responses:
      200:
        description: Массив дней с вложенными уроками
    """
    group_name = request.args.get('group', '')
    conn = get_db()
    query = 'SELECT * FROM schedule_days WHERE deleted_at IS NULL'
    params = []
    if group_name:
        query += ' AND group_name = ?'
        params.append(group_name)
    query += ' ORDER BY sort_order'
    days = conn.execute(query, params).fetchall()
    result = []
    for day in days:
        lessons = conn.execute(
            LESSON_SELECT + ' WHERE l.day_id = ? AND l.deleted_at IS NULL ORDER BY l.period',
            (day['id'],)
        ).fetchall()
        result.append({
            "id": day['id'],
            "day_name": day['day_name'],
            "date": day['date'],
            "sort_order": day['sort_order'],
            "group_name": day['group_name'],
            "lessons": [dict(l) for l in lessons]
        })
    conn.close()
    return jsonify(result)


@bp.route('/api/teacher-schedule', methods=['GET'])
def get_teacher_schedule():
    """Всё расписание преподавателя (все даты)
    ---
    parameters:
      - in: query
        name: teacher_id
        type: integer
        required: true
        description: ID преподавателя
    responses:
      200:
        description: Массив дней с уроками
    """
    teacher_id = request.args.get('teacher_id', '')
    if not teacher_id:
        return jsonify([])
    teacher_id = int(teacher_id)
    conn = get_db()
    lessons = conn.execute(
        LESSON_SELECT_WITH_DAY + ' WHERE l.teacher_id = ? AND l.subject_id IS NOT NULL AND l.deleted_at IS NULL ORDER BY sd.sort_order, l.period',
        (teacher_id,)
    ).fetchall()
    conn.close()

    days_map = {}
    for l in lessons:
        day_key = (l['day_name'], l['date'])
        if day_key not in days_map:
            days_map[day_key] = {
                'day_name': l['day_name'],
                'date': l['date'],
                'lessons': []
            }
        days_map[day_key]['lessons'].append({
            'id': l['id'],
            'period': l['period'],
            'subject': l['subject_name'],
            'teacher': l['teacher_name'],
            'classroom': l['classroom_name'],
            'groups': l['group_name'] or l['groups']
        })

    return jsonify(list(days_map.values()))


@bp.route('/api/teachers-with-schedule')
def teachers_with_schedule():
    """Список преподавателей, у которых есть уроки
    ---
    responses:
      200:
        description: Массив объектов {id, name}
    """
    conn = get_db()
    rows = conn.execute('''
        SELECT DISTINCT t.id, t.name
        FROM lessons l
        JOIN teachers t ON l.teacher_id = t.id
        JOIN schedule_days sd ON l.day_id = sd.id
        WHERE l.subject_id IS NOT NULL AND l.deleted_at IS NULL
        ORDER BY t.name
    ''').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name']} for r in rows])


@bp.route('/api/validate', methods=['POST'])
@login_required
def validate_lesson():
    """Проверить конфликты урока (учитель, аудитория, дубликат)
    ---
    parameters:
      - in: body
        name: body
        required: true
        schema:
          type: object
          properties:
            day_id:
              type: integer
            period:
              type: integer
            teacher_id:
              type: integer
            classroom_id:
              type: integer
            subject_id:
              type: integer
            group_id:
              type: integer
            exclude_lesson_id:
              type: integer
    responses:
      200:
        description: Объект с warnings и blocking
    """
    data = request.get_json()
    day_id = data.get('day_id')
    period = data.get('period')
    teacher_id = data.get('teacher_id')
    classroom_id = data.get('classroom_id')
    subject_id = data.get('subject_id')
    group_id = data.get('group_id')
    subgroup = data.get('subgroup')
    exclude_lesson_id = data.get('exclude_lesson_id')

    if not day_id or not period:
        return jsonify({"warnings": [], "blocking": False})

    conn = get_db()
    day = conn.execute('SELECT * FROM schedule_days WHERE id = ? AND deleted_at IS NULL', (day_id,)).fetchall()
    if not day:
        conn.close()
        return jsonify({"warnings": [], "blocking": False})
    day = day[0]

    warnings = []
    blocking = False

    if teacher_id:
        existing = conn.execute('''
            SELECT l.*, sd.day_name, sd.date, sd.group_name, t.name AS teacher_name,
                   CASE WHEN l.subgroup IS NOT NULL THEN gl.name || ' (' || l.subgroup || ')' ELSE gl.name END AS groups
            FROM lessons l
            JOIN schedule_days sd ON l.day_id = sd.id
            LEFT JOIN teachers t ON l.teacher_id = t.id
            LEFT JOIN groups_list gl ON l.group_id = gl.id
            WHERE l.teacher_id = ? AND l.period = ? AND sd.date = ?
            AND l.subject_id IS NOT NULL AND l.deleted_at IS NULL AND sd.deleted_at IS NULL
        ''', (teacher_id, period, day['date'])).fetchall()
        if exclude_lesson_id:
            existing = [e for e in existing if e['id'] != exclude_lesson_id]
        if existing:
            tname = existing[0]['teacher_name'] or ''
            groups = [e['group_name'] or e['groups'] for e in existing]
            warnings.append({
                "type": "teacher_conflict",
                "message": f"Преподаватель {tname} уже занят в {existing[0]['day_name']} на {period}-й паре (группа: {', '.join(groups)})"
            })
            blocking = True

    if classroom_id:
        existing = conn.execute('''
            SELECT l.*, sd.day_name, sd.date, sd.group_name, c.name AS classroom_name,
                   CASE WHEN l.subgroup IS NOT NULL THEN gl.name || ' (' || l.subgroup || ')' ELSE gl.name END AS groups
            FROM lessons l
            JOIN schedule_days sd ON l.day_id = sd.id
            LEFT JOIN classrooms_list c ON l.classroom_id = c.id
            LEFT JOIN groups_list gl ON l.group_id = gl.id
            WHERE l.classroom_id = ? AND l.period = ? AND sd.date = ?
            AND l.subject_id IS NOT NULL AND l.deleted_at IS NULL AND sd.deleted_at IS NULL
        ''', (classroom_id, period, day['date'])).fetchall()
        if exclude_lesson_id:
            existing = [e for e in existing if e['id'] != exclude_lesson_id]
        if existing:
            cname = existing[0]['classroom_name'] or ''
            groups = [e['group_name'] or e['groups'] for e in existing]
            warnings.append({
                "type": "classroom_conflict",
                "message": f"Аудитория {cname} уже занята в {existing[0]['day_name']} на {period}-й паре (группа: {', '.join(groups)})"
            })
            blocking = True

    if group_id is not None and subject_id:
        existing = conn.execute('''
            SELECT l.* FROM lessons l
            WHERE l.day_id = ? AND l.period = ? AND l.subject_id IS NOT NULL AND l.deleted_at IS NULL
        ''', (day_id, period)).fetchall()
        if exclude_lesson_id:
            existing = [e for e in existing if e['id'] != exclude_lesson_id]
        for e in existing:
            e_sub = e['subgroup']
            new_sub = subgroup
            e_gid = e['group_id']
            if e_gid == group_id and (e_sub or '') == (new_sub or ''):
                warnings.append({
                    "type": "duplicate",
                    "message": "Такой же урок уже добавлен в этот день на эту пару"
                })
                blocking = True
                break

    conn.close()
    return jsonify({"warnings": warnings, "blocking": blocking})


@bp.route('/api/week', methods=['DELETE'])
@login_required
def delete_week():
    """Удалить всю неделю (мягкое удаление)
    ---
    parameters:
      - in: query
        name: monday
        type: string
        required: true
        description: Дата понедельника (YYYY-MM-DD)
    responses:
      200:
        description: Объект с days_deleted
      400:
        description: monday обязателен
    """
    monday = request.args.get('monday', '')
    if not monday:
        return jsonify({"error": "monday required"}), 400
    week_dates = [d['date'] for d in build_week_dates(monday)]
    conn = get_db()
    try:
        conn.execute('BEGIN IMMEDIATE')
        placeholders = ','.join('?' * len(week_dates))
        conn.execute(
            f"UPDATE lessons SET deleted_at = datetime('now') WHERE deleted_at IS NULL AND day_id IN (SELECT id FROM schedule_days WHERE date IN ({placeholders}))",
            week_dates
        )
        conn.execute(
            f"UPDATE schedule_days SET deleted_at = datetime('now') WHERE deleted_at IS NULL AND date IN ({placeholders})",
            week_dates
        )
        affected = conn.execute(
            f"SELECT COUNT(*) FROM schedule_days WHERE date IN ({placeholders}) AND deleted_at IS NOT NULL",
            week_dates
        ).fetchone()[0]
        conn.commit()
    finally:
        conn.close()
    return jsonify({"success": True, "days_deleted": affected})


@bp.route('/api/week/restore', methods=['POST'])
@login_required
def restore_week():
    """Восстановить удалённую неделю
    ---
    parameters:
      - in: query
        name: monday
        type: string
        required: true
        description: Дата понедельника (YYYY-MM-DD)
    responses:
      200:
        description: Объект с days_restored
      400:
        description: monday обязателен
    """
    monday = request.args.get('monday', '')
    if not monday:
        return jsonify({"error": "monday required"}), 400
    week_dates = [d['date'] for d in build_week_dates(monday)]
    conn = get_db()
    try:
        conn.execute('BEGIN IMMEDIATE')
        placeholders = ','.join('?' * len(week_dates))
        conn.execute(
            f"UPDATE lessons SET deleted_at = NULL WHERE deleted_at IS NOT NULL AND day_id IN (SELECT id FROM schedule_days WHERE date IN ({placeholders}))",
            week_dates
        )
        conn.execute(
            f"UPDATE schedule_days SET deleted_at = NULL WHERE deleted_at IS NOT NULL AND date IN ({placeholders})",
            week_dates
        )
        affected = conn.execute(
            f"SELECT COUNT(*) FROM schedule_days WHERE date IN ({placeholders}) AND deleted_at IS NULL",
            week_dates
        ).fetchone()[0]
        conn.commit()
    finally:
        conn.close()
    return jsonify({"success": True, "days_restored": affected})


@bp.route('/api/validate-week', methods=['POST'])
@login_required
def validate_week():
    """Проверить, есть ли уже расписание на указанные даты
    ---
    parameters:
      - in: body
        name: body
        required: true
        schema:
          type: object
          properties:
            group_name:
              type: string
            dates:
              type: array
              items:
                type: string
    responses:
      200:
        description: Объект с ok (bool) и warnings
    """
    data = request.get_json()
    group_name = data.get('group_name', '')
    dates = data.get('dates', [])

    if not group_name or not dates:
        return jsonify({"ok": True, "warnings": []})

    conn = get_db()
    warnings = []

    placeholders = ','.join('?' * len(dates))
    existing = conn.execute(
        f'SELECT * FROM schedule_days WHERE group_name = ? AND deleted_at IS NULL AND date IN ({placeholders})',
        [group_name] + dates
    ).fetchall()

    if existing:
        existing_dates = [e['date'] for e in existing]
        warnings.append({
            "type": "week_exists",
            "message": f"У группы {group_name} уже есть расписание на даты: {', '.join(existing_dates)}. Удалите старую неделю перед добавлением новой."
        })

    conn.close()
    return jsonify({"ok": len(warnings) == 0, "warnings": warnings})


@bp.route('/api/days', methods=['POST'])
@login_required
def create_day():
    """Создать один день расписания
    ---
    parameters:
      - in: body
        name: body
        required: true
        schema:
          type: object
          properties:
            day_name:
              type: string
            date:
              type: string
            group_name:
              type: string
    responses:
      201:
        description: Созданный день
    """
    data = request.get_json()
    conn = get_db()
    max_order = conn.execute('SELECT COALESCE(MAX(sort_order), -1) FROM schedule_days').fetchone()[0]
    cursor = conn.execute(
        'INSERT INTO schedule_days (day_name, date, sort_order, group_name) VALUES (?, ?, ?, ?)',
        (data['day_name'], data['date'], max_order + 1, data.get('group_name', ''))
    )
    day_id = cursor.lastrowid
    conn.commit()
    conn.close()
    return jsonify({"id": day_id}), 201


@bp.route('/api/days/batch', methods=['POST'])
@login_required
def create_days_batch():
    """Создать неделю расписания (6 дней)
    ---
    parameters:
      - in: body
        name: body
        required: true
        schema:
          type: object
          properties:
            days:
              type: array
              items:
                type: object
                properties:
                  day_name:
                    type: string
                  date:
                    type: string
                  group_name:
                    type: string
    responses:
      201:
        description: Массив ids созданных дней
      400:
        description: Нет дней
      409:
        description: Уже есть расписание на эти даты
    """
    data = request.get_json()
    days = data.get('days', [])
    if not days:
        return jsonify({"error": "Нет дней"}), 400

    group_name = days[0].get('group_name', '')
    dates = [d['date'] for d in days]

    conn = get_db()

    placeholders = ','.join('?' * len(dates))
    existing = conn.execute(
        f'SELECT date FROM schedule_days WHERE group_name = ? AND deleted_at IS NULL AND date IN ({placeholders})',
        [group_name] + dates
    ).fetchall()

    if existing:
        existing_dates = [e['date'] for e in existing]
        conn.close()
        return jsonify({
            "error": f"У группы {group_name} уже есть расписание на даты: {', '.join(existing_dates)}"
        }), 409

    max_order = conn.execute('SELECT COALESCE(MAX(sort_order), -1) FROM schedule_days').fetchone()[0]
    ids = []
    for i, d in enumerate(days):
        cursor = conn.execute(
            'INSERT INTO schedule_days (day_name, date, sort_order, group_name) VALUES (?, ?, ?, ?)',
            (d['day_name'], d['date'], max_order + 1 + i, d.get('group_name', ''))
        )
        ids.append(cursor.lastrowid)
    conn.commit()
    conn.close()
    return jsonify({"ids": ids}), 201


@bp.route('/api/days/<int:day_id>', methods=['PUT'])
@login_required
def update_day(day_id):
    """Обновить день расписания
    ---
    parameters:
      - in: path
        name: day_id
        type: integer
        required: true
      - in: body
        name: body
        required: true
        schema:
          type: object
          properties:
            day_name:
              type: string
            date:
              type: string
    responses:
      200:
        description: success
    """
    data = request.get_json()
    conn = get_db()
    conn.execute(
        'UPDATE schedule_days SET day_name = ?, date = ? WHERE id = ?',
        (data['day_name'], data['date'], day_id)
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})


@bp.route('/api/days/<int:day_id>', methods=['DELETE'])
@login_required
def delete_day(day_id):
    """Мягкое удаление дня и его уроков
    ---
    parameters:
      - in: path
        name: day_id
        type: integer
        required: true
    responses:
      200:
        description: success
    """
    conn = get_db()
    conn.execute("UPDATE schedule_days SET deleted_at = datetime('now') WHERE id = ? AND deleted_at IS NULL", (day_id,))
    conn.execute("UPDATE lessons SET deleted_at = datetime('now') WHERE day_id = ?", (day_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})


@bp.route('/api/days/<int:day_id>/restore', methods=['POST'])
@login_required
def restore_day(day_id):
    """Восстановить мягко удалённый день
    ---
    parameters:
      - in: path
        name: day_id
        type: integer
        required: true
    responses:
      200:
        description: success
      409:
        description: Конфликт при восстановлении
    """
    conn = get_db()
    day = conn.execute("SELECT id, day_name, date, group_name FROM schedule_days WHERE id = ?", (day_id,)).fetchone()
    if not day:
        conn.close()
        return jsonify({"error": "Day not found"}), 404

    try:
        conn.execute('BEGIN IMMEDIATE')
        lessons = conn.execute(
            "SELECT id, day_id, period, teacher_id, classroom_id, subject_id, group_id, subgroup FROM lessons WHERE day_id = ? AND deleted_at IS NOT NULL",
            (day_id,)
        ).fetchall()

        conflict_warnings = []
        for l in lessons:
            ws, blocking = _validate_lesson_data(conn, l['day_id'], l['period'], l['teacher_id'], l['classroom_id'],
                                                  subject_id=l['subject_id'], group_id=l['group_id'], subgroup=l['subgroup'],
                                                  skip_day_check=True)
            if blocking:
                conflict_warnings.extend(ws)

        if conflict_warnings:
            conn.rollback()
            conn.close()
            return jsonify({"error": "Конфликт при восстановлении дня", "warnings": conflict_warnings}), 409

        conn.execute("UPDATE schedule_days SET deleted_at = NULL WHERE id = ?", (day_id,))
        conn.execute("UPDATE lessons SET deleted_at = NULL WHERE day_id = ? AND deleted_at IS NOT NULL", (day_id,))
        conn.commit()
    finally:
        conn.close()

    return jsonify({
        "success": True,
        "item": {
            "type": "day",
            "day_name": day['day_name'],
            "date": day['date'],
            "group_name": day['group_name'],
        }
    })


@bp.route('/api/days/<int:day_id>/lessons', methods=['POST'])
@login_required
def create_lesson(day_id):
    """Создать урок в дне расписания
    ---
    parameters:
      - in: path
        name: day_id
        type: integer
        required: true
      - in: body
        name: body
        required: true
        schema:
          type: object
          properties:
            period:
              type: integer
            subject_id:
              type: integer
            teacher_id:
              type: integer
            classroom_id:
              type: integer
            group_id:
              type: integer
            subgroup:
              type: integer
    responses:
      201:
        description: Созданный урок
      400:
        description: Invalid period
      409:
        description: Conflict detected
    """
    data = request.get_json()
    conn = get_db()
    try:
        conn.execute('BEGIN IMMEDIATE')

        teacher_id = data.get('teacher_id')
        classroom_id = data.get('classroom_id')
        subject_id = data.get('subject_id')
        period = data.get('period')
        if not period or period not in range(1, 9):
            conn.rollback()
            return jsonify({"error": "Invalid period (must be 1-8)"}), 400
        group_id = data.get('group_id')
        subgroup = data.get('subgroup')

        warnings, blocking = _validate_lesson_data(conn, day_id, period, teacher_id, classroom_id, subject_id, group_id=group_id, subgroup=subgroup)
        if blocking:
            conn.rollback()
            return jsonify({"error": "Conflict detected", "warnings": warnings}), 409

        cursor = conn.execute(
            'INSERT INTO lessons (day_id, period, subject_id, teacher_id, classroom_id, group_id, subgroup) VALUES (?, ?, ?, ?, ?, ?, ?)',
            (day_id, period, subject_id, teacher_id, classroom_id, group_id, subgroup)
        )
        lesson_id = cursor.lastrowid
        conn.commit()
        return jsonify({"id": lesson_id}), 201
    finally:
        conn.close()


@bp.route('/api/lessons/<int:lesson_id>', methods=['PUT'])
@login_required
def update_lesson(lesson_id):
    """Обновить урок (сохраняет историю изменений)
    ---
    parameters:
      - in: path
        name: lesson_id
        type: integer
        required: true
      - in: body
        name: body
        required: true
        schema:
          type: object
          properties:
            subject_id:
              type: integer
            teacher_id:
              type: integer
            classroom_id:
              type: integer
            group_id:
              type: integer
    responses:
      200:
        description: success
      404:
        description: Lesson not found
      409:
        description: Conflict detected
    """
    data = request.get_json()
    conn = get_db()
    try:
        conn.execute('BEGIN IMMEDIATE')

        existing = conn.execute('SELECT * FROM lessons WHERE id = ? AND deleted_at IS NULL', (lesson_id,)).fetchone()
        if not existing:
            conn.rollback()
            return jsonify({"error": "Lesson not found"}), 404

        teacher_id = data.get('teacher_id')
        classroom_id = data.get('classroom_id')
        subject_id = data.get('subject_id')
        day_id = existing['day_id']
        period = data.get('period', existing['period'])
        group_id = data.get('group_id')
        subgroup = data.get('subgroup')

        warnings, blocking = _validate_lesson_data(conn, day_id, period, teacher_id, classroom_id, subject_id, exclude_id=lesson_id, group_id=group_id, subgroup=subgroup)
        if blocking:
            conn.rollback()
            return jsonify({"error": "Conflict detected", "warnings": warnings}), 409

        conn.execute(
            'INSERT INTO lesson_history (lesson_id, old_subject_id, old_teacher_id, old_classroom_id, old_group_id, old_subgroup) VALUES (?, ?, ?, ?, ?, ?)',
            (lesson_id, existing['subject_id'], existing['teacher_id'], existing['classroom_id'], existing['group_id'], existing['subgroup'])
        )

        conn.execute(
            'UPDATE lessons SET subject_id = ?, teacher_id = ?, classroom_id = ?, group_id = ?, subgroup = ? WHERE id = ?',
            (subject_id, teacher_id, classroom_id, group_id, subgroup, lesson_id)
        )
        conn.commit()
        return jsonify({"success": True})
    finally:
        conn.close()


@bp.route('/api/lessons/<int:lesson_id>', methods=['DELETE'])
@login_required
def delete_lesson(lesson_id):
    """Мягкое удаление урока
    ---
    parameters:
      - in: path
        name: lesson_id
        type: integer
        required: true
    responses:
      200:
        description: success
    """
    conn = get_db()
    conn.execute("UPDATE lessons SET deleted_at = datetime('now') WHERE id = ? AND deleted_at IS NULL", (lesson_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})


@bp.route('/api/lessons/<int:lesson_id>/restore', methods=['POST'])
@login_required
def restore_lesson(lesson_id):
    """Восстановить мягко удалённый урок
    ---
    parameters:
      - in: path
        name: lesson_id
        type: integer
        required: true
    responses:
      200:
        description: success
      409:
        description: Конфликт при восстановлении
    """
    conn = get_db()
    lesson = conn.execute(LESSON_SELECT_WITH_DAY + ' WHERE l.id = ?', (lesson_id,)).fetchone()
    if not lesson:
        conn.close()
        return jsonify({"error": "Lesson not found"}), 404

    day_id = lesson['day_id']
    period = lesson['period']
    teacher_id = lesson['teacher_id']
    classroom_id = lesson['classroom_id']
    subject_id = lesson['subject_id']
    group_id = lesson['group_id']
    subgroup = lesson['subgroup']

    warnings, blocking = _validate_lesson_data(conn, day_id, period, teacher_id, classroom_id, subject_id=subject_id, group_id=group_id, subgroup=subgroup)
    if blocking:
        conn.close()
        return jsonify({"error": "Конфликт при восстановлении", "warnings": warnings}), 409

    conn.execute("UPDATE lessons SET deleted_at = NULL WHERE id = ?", (lesson_id,))
    conn.commit()
    conn.close()
    return jsonify({
        "success": True,
        "item": {
            "type": "lesson",
            "subject_name": lesson['subject_name'],
            "teacher_name": lesson['teacher_name'],
            "day_name": lesson['day_name'],
            "date": lesson['date'],
            "period": lesson['period'],
            "group_name": lesson['group_name'],
        }
    })


@bp.route('/api/trash')
@login_required
def get_trash():
    """Корзина — мягко удалённые дни и уроки за последние {TRASH_RETENTION_DAYS} дней
    ---
    responses:
      200:
        description: Объект с массивами days и lessons
    """
    conn = get_db()
    retention_param = f'-{TRASH_RETENTION_DAYS} days'
    deleted_days = conn.execute('''
        SELECT d.id, d.day_name, d.date, d.group_name, d.deleted_at,
               'День' AS type_label
        FROM schedule_days d
        WHERE d.deleted_at IS NOT NULL AND d.deleted_at >= datetime('now', ?)
        ORDER BY d.deleted_at DESC
    ''', (retention_param,)).fetchall()
    deleted_lessons = conn.execute(
        "SELECT l.id, l.period, l.deleted_at, sd.day_name, sd.date, sd.group_name,"
        " s.name AS subject_name, t.name AS teacher_name, c.name AS classroom_name,"
        " CASE WHEN l.subgroup IS NOT NULL THEN gl.name || ' (' || l.subgroup || ')' ELSE gl.name END AS groups,"
        " 'Урок' AS type_label"
        " FROM lessons l"
        " JOIN schedule_days sd ON l.day_id = sd.id"
        " LEFT JOIN teachers t ON l.teacher_id = t.id"
        " LEFT JOIN subjects s ON l.subject_id = s.id"
        " LEFT JOIN classrooms_list c ON l.classroom_id = c.id"
        " LEFT JOIN groups_list gl ON l.group_id = gl.id"
        " WHERE l.deleted_at IS NOT NULL AND l.deleted_at >= datetime('now', ?) ORDER BY l.deleted_at DESC",
        (retention_param,)
    ).fetchall()
    deleted_refs = {}
    type_labels = {'subjects': 'Предмет', 'teachers': 'Преподаватель', 'classrooms': 'Аудитория', 'groups': 'Группа'}
    for table, label in [('subjects', 'subjects'), ('teachers', 'teachers'), ('classrooms_list', 'classrooms'), ('groups_list', 'groups')]:
        rows = conn.execute(f"SELECT id, name, deleted_at, '{type_labels[label]}' AS type_label FROM {table} WHERE deleted_at IS NOT NULL AND deleted_at >= datetime('now', ?) ORDER BY deleted_at DESC", (retention_param,)).fetchall()
        deleted_refs[label] = [dict(r) for r in rows]
    conn.close()
    return jsonify({
        "days": [dict(d) for d in deleted_days],
        "lessons": [dict(l) for l in deleted_lessons],
        **deleted_refs,
    })


@bp.route('/api/classrooms-with-schedule')
def classrooms_with_schedule():
    """Список аудиторий, в которых есть занятия
    ---
    responses:
      200:
        description: Массив объектов {id, name}
    """
    conn = get_db()
    rows = conn.execute('''
        SELECT DISTINCT c.id, c.name
        FROM lessons l
        JOIN classrooms_list c ON l.classroom_id = c.id
        JOIN schedule_days sd ON l.day_id = sd.id
        WHERE l.subject_id IS NOT NULL AND l.deleted_at IS NULL
        ORDER BY CAST(c.name AS INTEGER), c.name
    ''').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name']} for r in rows])


@bp.route('/api/schedule-table')
def schedule_table():
    """Расписание группы на неделю
    ---
    parameters:
      - in: query
        name: group
        type: string
        required: true
        description: Название группы
      - in: query
        name: monday
        type: string
        required: true
        description: Дата понедельника (YYYY-MM-DD)
    responses:
      200:
        description: Объект с days (массив дней) и periods (массив пар с уроками)
    """
    group_name = request.args.get('group', '')
    monday = request.args.get('monday', '')
    if not group_name or not monday:
        return jsonify({'days': [], 'periods': []})

    week_dates, lessons_by_date = _fetch_group_raw(group_name, monday)

    periods = []
    for b in BELL_SCHEDULE:
        p = b['period']
        lessons = {}
        for wd in week_dates:
            d = wd['date']
            if d in lessons_by_date and p in lessons_by_date[d]:
                raw_list = lessons_by_date[d][p]
                lessons[d] = [{
                    'id': l['id'], 'subject': l['subject_name'], 'teacher': l['teacher_name'],
                    'classroom': l['classroom_name'], 'groups': l['groups'],
                    'day_id': l['day_id'],
                    'subject_id': l['subject_id'], 'teacher_id': l['teacher_id'],
                    'classroom_id': l['classroom_id'],
                    'group_id': l['group_id'], 'subgroup': l['subgroup']
                } for l in raw_list]
            else:
                lessons[d] = []
        periods.append({**b, 'lessons': lessons})

    return jsonify({'days': week_dates, 'periods': periods, 'week_exists': bool(lessons_by_date)})


@bp.route('/api/teacher-schedule-table')
def teacher_schedule_table():
    """Расписание преподавателя на неделю
    ---
    parameters:
      - in: query
        name: teacher_id
        type: integer
        required: true
        description: ID преподавателя
      - in: query
        name: monday
        type: string
        required: true
        description: Дата понедельника (YYYY-MM-DD)
    responses:
      200:
        description: Объект с days и periods
    """
    teacher_id = request.args.get('teacher_id', '')
    monday = request.args.get('monday', '')
    if not teacher_id or not monday:
        return jsonify({'days': [], 'periods': []})
    teacher_id = int(teacher_id)

    week_dates, lessons_by_date_period = _fetch_entity_raw('teachers', teacher_id, 'teacher_id', monday)

    return jsonify({'days': week_dates, 'periods': _build_schedule_periods(week_dates, lessons_by_date_period)})


@bp.route('/api/classroom-schedule-table')
def classroom_schedule_table():
    """Занятость аудитории на неделю
    ---
    parameters:
      - in: query
        name: classroom_id
        type: integer
        required: true
        description: ID аудитории
      - in: query
        name: monday
        type: string
        required: true
        description: Дата понедельника (YYYY-MM-DD)
    responses:
      200:
        description: Объект с days и periods
    """
    classroom_id = request.args.get('classroom_id', '')
    monday = request.args.get('monday', '')
    if not classroom_id or not monday:
        return jsonify({'days': [], 'periods': []})
    classroom_id = int(classroom_id)

    week_dates, lessons_by_date_period = _fetch_entity_raw('classrooms_list', classroom_id, 'classroom_id', monday)

    return jsonify({'days': week_dates, 'periods': _build_schedule_periods(week_dates, lessons_by_date_period)})


generic_crud('subjects', 'subjects', 'name', extra_fields=['default_teacher_id'])
generic_crud('teachers', 'teachers', 'name', extra_fields=['phone'])
generic_crud('classrooms_list', 'classrooms', 'CAST(name AS INTEGER), name')


def _make_restore(table, label):
    @login_required
    def restore_item(item_id):
        conn = get_db()
        item = conn.execute(f"SELECT id, name FROM {table} WHERE id = ?", (item_id,)).fetchone()
        if not item:
            conn.close()
            return jsonify({"error": "Not found"}), 404
        if table == 'teachers':
            phone = conn.execute("SELECT phone FROM teachers WHERE id = ? AND deleted_at IS NOT NULL", (item_id,)).fetchone()
            if phone and phone['phone'] and phone['phone'].strip():
                existing = conn.execute("SELECT 1 FROM teachers WHERE phone = ? AND id != ? AND deleted_at IS NULL", (phone['phone'].strip(), item_id)).fetchone()
                if existing:
                    conn.close()
                    return jsonify({"error": "Этот номер телефона уже используется другим преподавателем"}), 409
        conn.execute(f"UPDATE {table} SET deleted_at = NULL WHERE id = ?", (item_id,))
        conn.commit()
        conn.close()
        refs_cache.invalidate_prefix(f'ref_list_{table}')
        return jsonify({"success": True, "item": {"type": label, "name": item['name']}})
    restore_item.__name__ = f'restore_{label}'
    bp.add_url_rule(f'/api/{label}/<int:item_id>/restore', f'restore_{label}', restore_item, methods=['POST'])


_make_restore('subjects', 'subjects')
_make_restore('teachers', 'teachers')
_make_restore('classrooms_list', 'classrooms')


@bp.route('/api/groups', methods=['GET'])
def get_groups():
    """Список групп
    ---
    responses:
      200:
        description: Массив групп с полями id, name, course, is_distance
    """
    cached = refs_cache.get('ref_list_groups_list')
    if cached is not None:
        return jsonify(cached)
    conn = get_db()
    items = conn.execute("SELECT * FROM groups_list WHERE deleted_at IS NULL ORDER BY course, name").fetchall()
    conn.close()
    result = [dict(i) for i in items]
    refs_cache.set('ref_list_groups_list', result)
    return jsonify(result)


@bp.route('/api/groups', methods=['POST'])
@login_required
def create_group():
    """Создать группу
    ---
    parameters:
      - in: body
        name: body
        required: true
        schema:
          type: object
          properties:
            name:
              type: string
            course:
              type: integer
            is_distance:
              type: boolean
    responses:
      201:
        description: Созданная группа
      400:
        description: Name required или группа уже существует
    """
    data = request.get_json()
    name = data.get('name', '').strip()
    if not name:
        return jsonify({"error": "Name required"}), 400
    course = data.get('course', 1)
    is_distance = 1 if data.get('is_distance') else 0
    conn = get_db()
    active = conn.execute("SELECT id FROM groups_list WHERE LOWER(name) = LOWER(?) AND deleted_at IS NULL", (name,)).fetchone()
    if active:
        conn.close()
        return jsonify({"error": "Группа с таким названием уже существует"}), 400
    deleted = conn.execute("SELECT id, course, is_distance FROM groups_list WHERE LOWER(name) = LOWER(?) AND deleted_at IS NOT NULL", (name,)).fetchone()
    if deleted:
        conn.execute("UPDATE groups_list SET deleted_at = NULL, course = ?, is_distance = ? WHERE id = ?", (course, is_distance, deleted['id']))
        conn.commit()
        conn.close()
        refs_cache.invalidate_prefix('ref_list_groups_list')
        return jsonify({"id": deleted['id'], "name": name, "course": course, "is_distance": is_distance, "restored": True}), 200
    try:
        cursor = conn.execute(
            'INSERT INTO groups_list (name, course, is_distance) VALUES (?, ?, ?)',
            (name, course, is_distance)
        )
        conn.commit()
        item_id = cursor.lastrowid
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": "Уже существует"}), 409
    conn.close()
    refs_cache.invalidate_prefix('ref_list_groups_list')
    return jsonify({"id": item_id, "name": name, "course": course, "is_distance": is_distance}), 201


@bp.route('/api/groups/<int:item_id>', methods=['PUT'])
@login_required
def update_group(item_id):
    """Обновить группу
    ---
    parameters:
      - in: path
        name: item_id
        type: integer
        required: true
      - in: body
        name: body
        required: true
        schema:
          type: object
          properties:
            name:
              type: string
            course:
              type: integer
            is_distance:
              type: boolean
    responses:
      200:
        description: success
      400:
        description: Name required или дубликат
    """
    data = request.get_json()
    name = data.get('name', '').strip()
    if not name:
        return jsonify({"error": "Name required"}), 400
    course = data.get('course', 1)
    is_distance = 1 if data.get('is_distance') else 0
    conn = get_db()
    existing = conn.execute(
        "SELECT id FROM groups_list WHERE LOWER(name) = LOWER(?) AND id != ?",
        (name, item_id)
    ).fetchone()
    if existing:
        conn.close()
        return jsonify({"error": "Группа с таким названием уже существует"}), 400
    conn.execute(
        'UPDATE groups_list SET name = ?, course = ?, is_distance = ? WHERE id = ?',
        (name, course, is_distance, item_id)
    )
    conn.commit()
    conn.close()
    refs_cache.invalidate_prefix('ref_list_groups_list')
    return jsonify({"success": True})


@bp.route('/api/groups/<int:item_id>', methods=['DELETE'])
@login_required
def delete_group(item_id):
    """Удалить группу
    ---
    parameters:
      - in: path
        name: item_id
        type: integer
        required: true
    responses:
      200:
        description: success
    """
    conn = get_db()
    conn.execute("UPDATE groups_list SET deleted_at = datetime('now') WHERE id = ? AND deleted_at IS NULL", (item_id,))
    conn.commit()
    conn.close()
    refs_cache.invalidate_prefix('ref_list_groups_list')
    return jsonify({"success": True, "warning": "Запись помечена как удалённая."})


@bp.route('/api/groups/<int:item_id>/restore', methods=['POST'])
@login_required
def restore_group(item_id):
    conn = get_db()
    item = conn.execute("SELECT id, name FROM groups_list WHERE id = ?", (item_id,)).fetchone()
    if not item:
        conn.close()
        return jsonify({"error": "Not found"}), 404
    conn.execute("UPDATE groups_list SET deleted_at = NULL WHERE id = ?", (item_id,))
    conn.commit()
    conn.close()
    refs_cache.invalidate_prefix('ref_list_groups_list')
    return jsonify({"success": True, "item": {"type": "group", "name": item['name']}})


@bp.route('/api/groups/advance-courses', methods=['POST'])
@login_required
def advance_courses():
    """Перевести все группы на следующий курс (+1), удалить 4-й курс
    ---
    responses:
      200:
        description: Объект с полем deleted (количество удалённых групп)
    """
    conn = get_db()
    deleted = conn.execute('DELETE FROM groups_list WHERE course = 4').rowcount
    conn.execute('UPDATE groups_list SET course = course + 1 WHERE course < 4')
    conn.commit()
    conn.close()
    refs_cache.invalidate_prefix('ref_list_groups_list')
    return jsonify({"success": True, "deleted": deleted})


@bp.route('/api/export')
def export_schedule():
    """Экспорт расписания в файл (xlsx/docx/pdf)
    ---
    parameters:
      - in: query
        name: type
        type: string
        enum: [group, teacher, classroom]
        required: true
        description: Тип расписания
      - in: query
        name: name
        type: string
        required: true
        description: Название группы / имя преподавателя / номер аудитории
      - in: query
        name: id
        type: integer
        required: false
        description: ID преподавателя/аудитории (для teacher/classroom)
      - in: query
        name: monday
        type: string
        required: true
        description: Дата понедельника (YYYY-MM-DD)
      - in: query
        name: format
        type: string
        enum: [xlsx, docx, pdf]
        required: true
        description: Формат файла
    responses:
      200:
        description: Файл расписания
      400:
        description: Отсутствуют параметры или неверный формат
      404:
        description: Нет данных для экспорта
    """
    schedule_type = request.args.get('type', '')
    name = request.args.get('name', '')
    entity_id = request.args.get('id', '')
    monday = request.args.get('monday', '')
    fmt = request.args.get('format', '')

    if not all([schedule_type, name, monday, fmt]):
        return jsonify({"error": "Missing parameters"}), 400
    if fmt not in ('xlsx', 'docx', 'pdf'):
        return jsonify({"error": "Invalid format"}), 400

    if schedule_type in ('teacher', 'classroom') and entity_id:
        try:
            entity_id = int(entity_id)
        except (ValueError, TypeError):
            entity_id = None

    week_dates, periods = _get_export_data(schedule_type, name, monday, entity_id)

    has_data = any(
        p['lessons'].get(d['date'], '') if isinstance(p['lessons'].get(d['date'], ''), str)
        else len(p['lessons'].get(d['date'], [])) > 0
        for p in periods for d in week_dates
    )
    logger.info("Export %s '%s'(id=%s) week %s: periods=%d, has_data=%s",
                schedule_type, name, entity_id, monday, len(periods), has_data)
    if not has_data:
        return jsonify({"error": "No schedule data"}), 404

    sat = datetime.strptime(week_dates[-1]['date'], '%Y-%m-%d')
    mon = datetime.strptime(week_dates[0]['date'], '%Y-%m-%d')
    date_range = f"{mon.day:02d}.{mon.month:02d}–{sat.day:02d}.{sat.month:02d}"

    type_labels = {'group': 'Группа', 'teacher': 'Преподаватель', 'classroom': 'Аудитория'}
    title = f"Расписание — {type_labels.get(schedule_type, '')} {name}, неделя {date_range}"
    safe_name = _safe_filename(f"raspisanie_{name}_{date_range}")

    if fmt == 'xlsx':
        buf = _export_excel(week_dates, periods, title)
        return send_file(buf, as_attachment=True, download_name=f"{safe_name}.xlsx",
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    elif fmt == 'docx':
        buf = _export_word(week_dates, periods, title)
        return send_file(buf, as_attachment=True, download_name=f"{safe_name}.docx",
                         mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    elif fmt == 'pdf':
        buf = _export_pdf(week_dates, periods, title)
        return send_file(buf, as_attachment=True, download_name=f"{safe_name}.pdf",
                         mimetype='application/pdf')


@bp.route('/api/availability')
def availability():
    """Свободные окна преподавателей и аудиторий
    ---
    parameters:
      - in: query
        name: monday
        type: string
        required: true
        description: Дата понедельника (YYYY-MM-DD)
      - in: query
        name: group_id
        type: integer
        required: false
        description: Фильтр по группе
    responses:
      200:
        description: Объект с days, periods, teachers, classrooms, teacher_busy, classroom_busy
      400:
        description: monday обязателен
    """
    monday = request.args.get('monday', '')
    group_id_filter = request.args.get('group_id', type=int)
    if not monday:
        return jsonify({"error": "monday required"}), 400

    week_dates = build_week_dates(monday)
    date_list = [wd['date'] for wd in week_dates]
    placeholders = ','.join('?' * len(date_list))

    conn = get_db()
    try:
        query = f'''
            SELECT l.period, sd.date, sd.day_name,
                   t.name AS teacher_name, l.teacher_id,
                   c.name AS classroom_name, l.classroom_id,
                   g.name AS group_name, l.group_id, g.course,
                   s.name AS subject_name, l.subject_id
            FROM lessons l
            JOIN schedule_days sd ON l.day_id = sd.id
            LEFT JOIN teachers t ON l.teacher_id = t.id AND t.deleted_at IS NULL
            LEFT JOIN classrooms_list c ON l.classroom_id = c.id AND c.deleted_at IS NULL
            LEFT JOIN groups_list g ON l.group_id = g.id
            LEFT JOIN subjects s ON l.subject_id = s.id
            WHERE sd.date IN ({placeholders})
              AND l.subject_id IS NOT NULL
              AND l.deleted_at IS NULL AND sd.deleted_at IS NULL
        '''
        params = date_list[:]
        if group_id_filter:
            query += ' AND l.group_id = ?'
            params.append(group_id_filter)
        lessons = conn.execute(query, params).fetchall()

        all_teachers = conn.execute('SELECT id, name FROM teachers WHERE deleted_at IS NULL ORDER BY name').fetchall()
        all_classrooms = conn.execute('SELECT id, name FROM classrooms_list WHERE deleted_at IS NULL ORDER BY CAST(name AS INTEGER), name').fetchall()
        all_groups = conn.execute('SELECT id, name, course FROM groups_list WHERE deleted_at IS NULL ORDER BY course, name').fetchall()

        teacher_busy = {}
        classroom_busy = {}
        teacher_courses = {}
        all_courses_set = set()

        for l in lessons:
            key = (l['date'], l['period'])
            t_id = l['teacher_id']
            c_id = l['classroom_id']
            course = l['course']
            info = {'group': l['group_name'], 'subject': l['subject_name'], 'teacher': l['teacher_name'], 'classroom': l['classroom_name'], 'teacher_id': t_id, 'classroom_id': c_id}
            if t_id and l['teacher_name']:
                teacher_busy.setdefault(t_id, {})[key] = info
                teacher_courses.setdefault(t_id, set()).add(course)
                all_courses_set.add(course)
            if c_id and l['classroom_name']:
                classroom_busy.setdefault(c_id, {})[key] = info

        periods = [{'period': b['period'], 'start': b['start'], 'end': b['end']} for b in BELL_SCHEDULE]

        teachers_out = []
        for t in all_teachers:
            td = {'id': t['id'], 'name': t['name']}
            courses = sorted(teacher_courses.get(t['id'], []))
            if courses:
                td['courses'] = courses
            teachers_out.append(td)

        return jsonify({
            'days': week_dates,
            'periods': periods,
            'teachers': teachers_out,
            'classrooms': [{'id': c['id'], 'name': c['name']} for c in all_classrooms],
            'groups': [{'id': g['id'], 'name': g['name'], 'course': g['course']} for g in all_groups],
            'all_courses': sorted(all_courses_set),
            'teacher_busy': {str(k): {f"{dk[0]}_{dk[1]}": v for dk, v in dv.items()} for k, dv in teacher_busy.items()},
            'classroom_busy': {str(k): {f"{dk[0]}_{dk[1]}": v for dk, v in dv.items()} for k, dv in classroom_busy.items()},
        })
    finally:
        conn.close()





@bp.route('/api/db/health')
@login_required
def db_health():
    conn = get_db()
    integrity = conn.execute("PRAGMA integrity_check").fetchone()[0]
    tables_info = {}
    with_deleted = {'schedule_days', 'lessons', 'subjects', 'teachers', 'groups_list', 'classrooms_list'}
    for t in ['schedule_days', 'lessons', 'subjects', 'teachers', 'groups_list', 'classrooms_list', 'excluded_dates', 'lesson_history']:
        total = conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
        d = conn.execute(f"SELECT COUNT(*) FROM {t} WHERE deleted_at IS NOT NULL").fetchone()[0] if t in with_deleted else 0
        tables_info[t] = {"total": total, "deleted": d}
    conflicts = conn.execute('''
        SELECT l1.id AS id1, l2.id AS id2, sd.date, l1.period,
               t1.name AS t1n, t2.name AS t2n
        FROM lessons l1 JOIN lessons l2 ON l1.id < l2.id
        JOIN schedule_days sd ON l1.day_id = sd.id
        LEFT JOIN teachers t1 ON l1.teacher_id = t1.id
        LEFT JOIN teachers t2 ON l2.teacher_id = t2.id
        WHERE l1.day_id = l2.day_id AND l1.period = l2.period
          AND l1.deleted_at IS NULL AND l2.deleted_at IS NULL
          AND (l1.teacher_id = l2.teacher_id OR l1.classroom_id = l2.classroom_id
               OR (l1.group_id = l2.group_id AND (l1.subgroup IS NULL OR l1.subgroup = l2.subgroup)))
    ''').fetchall()
    orphaned = []
    for t, col in [('subjects','subject_id'),('teachers','teacher_id'),('classrooms_list','classroom_id'),('groups_list','group_id')]:
        rows = conn.execute(f"SELECT l.id FROM lessons l LEFT JOIN {t} r ON l.{col}=r.id WHERE l.{col} IS NOT NULL AND r.id IS NULL AND l.deleted_at IS NULL").fetchall()
        orphaned.extend([dict(r) for r in rows])
    conn.close()
    return jsonify({"integrity_check": integrity, "tables": tables_info, "conflicts": [dict(r) for r in conflicts], "orphaned_lessons": orphaned})


@bp.route('/api/db/purge-trash', methods=['POST'])
@login_required
def purge_trash():
    conn = get_db()
    result = {}
    for t in ['lessons', 'schedule_days', 'subjects', 'teachers', 'groups_list', 'classrooms_list']:
        cnt = conn.execute(f"DELETE FROM {t} WHERE deleted_at IS NOT NULL AND deleted_at < datetime('now', ?)", (f'-{TRASH_RETENTION_DAYS} days',)).rowcount
        if cnt:
            result[t] = cnt
    conn.execute("DELETE FROM lesson_history WHERE changed_at < datetime('now', ?)", (f'-{TRASH_RETENTION_DAYS} days',))
    conn.commit()
    conn.close()
    refs_cache.invalidate_prefix('ref_list_')
    return jsonify({"purged": result})
