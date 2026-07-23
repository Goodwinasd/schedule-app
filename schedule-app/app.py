import sqlite3
import os
import re
import time
import io
import shutil
from datetime import datetime, timedelta
from flask import Flask, jsonify, request, render_template, session, send_file
from flask_cors import CORS
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash
import secrets


def _load_env():
    env_path = os.path.join(os.path.dirname(__file__), '.env')
    if os.path.exists(env_path):
        with open(env_path, encoding='utf-8') as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith('#') and '=' in line:
                    k, v = line.split('=', 1)
                    os.environ.setdefault(k.strip(), v.strip())

_load_env()

app = Flask(__name__)
_sk = os.environ.get('SCHEDULE_SECRET_KEY')
if not _sk:
    raise Exception("Не найдена переменная SCHEDULE_SECRET_KEY. Создайте файл .env рядом с app.py:\n  SCHEDULE_SECRET_KEY=случайная_строка_32_символа")
app.secret_key = _sk
app.permanent_session_lifetime = 3600
CORS(app)


@app.after_request
def set_charset(response):
    if response.content_type and 'text/html' in response.content_type:
        response.content_type = 'text/html; charset=utf-8'
    elif response.content_type and 'application/json' in response.content_type:
        response.content_type = 'application/json; charset=utf-8'
    if request.path.startswith('/static/'):
        response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
    return response


DB_PATH = os.path.join(os.path.dirname(__file__), 'schedule.db')

BELL_SCHEDULE = [
    {"period": 1, "start": "8:30", "end": "9:55", "shift": 1},
    {"period": 2, "start": "10:00", "end": "11:25", "shift": 1},
    {"period": 3, "start": "12:00", "end": "13:25", "shift": 1},
    {"period": 4, "start": "13:30", "end": "14:55", "shift": 2},
    {"period": 5, "start": "15:00", "end": "16:25", "shift": 2},
    {"period": 6, "start": "16:55", "end": "18:20", "shift": 2},
    {"period": 7, "start": "18:25", "end": "19:50", "shift": 3},
    {"period": 8, "start": "19:55", "end": "21:20", "shift": 3},
]

SHIFT_LABELS = {1: '1 смена', 2: '2 смена', 3: '3 смена'}

_pw = os.environ.get('SCHEDULE_PASSWORD')
if not _pw:
    raise Exception("Не найдена переменная SCHEDULE_PASSWORD. Создайте файл .env рядом с app.py:\n  SCHEDULE_PASSWORD=ваш_пароль")
ADMIN_PASSWORD_HASH = generate_password_hash(_pw)

_login_attempts = {}
LOGIN_RATE_LIMIT = 5
LOGIN_RATE_WINDOW = 300
_LOGIN_CLEANUP_INTERVAL = 300
_last_cleanup_time = 0


def _clean_login_attempts():
    global _last_cleanup_time
    now = time.time()
    if now - _last_cleanup_time < _LOGIN_CLEANUP_INTERVAL:
        return
    _last_cleanup_time = now
    expired = [ip for ip, ts in _login_attempts.items() if not ts or now - ts[-1] >= LOGIN_RATE_WINDOW]
    for ip in expired:
        _login_attempts.pop(ip, None)


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode=WAL")
    return conn


def init_db():
    conn = get_db()
    conn.executescript('''
        CREATE TABLE IF NOT EXISTS schedule_days (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            day_name TEXT NOT NULL,
            date TEXT NOT NULL,
            sort_order INTEGER NOT NULL DEFAULT 0,
            group_name TEXT DEFAULT '',
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS lessons (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            day_id INTEGER NOT NULL,
            period INTEGER NOT NULL,
            subject_id INTEGER,
            teacher_id INTEGER,
            classroom_id INTEGER,
            group_id INTEGER,
            subgroup INTEGER DEFAULT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (day_id) REFERENCES schedule_days(id) ON DELETE CASCADE,
            FOREIGN KEY (teacher_id) REFERENCES teachers(id),
            FOREIGN KEY (subject_id) REFERENCES subjects(id),
            FOREIGN KEY (classroom_id) REFERENCES classrooms_list(id),
            FOREIGN KEY (group_id) REFERENCES groups_list(id)
        );

        CREATE TABLE IF NOT EXISTS subjects (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE
        );

        CREATE TABLE IF NOT EXISTS teachers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS groups_list (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            course INTEGER DEFAULT 1,
            is_distance INTEGER DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS classrooms_list (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL
        );
    ''')
    conn.commit()

    # Migration: add course/is_distance to groups_list if missing
    cols = [r[1] for r in conn.execute("PRAGMA table_info(groups_list)").fetchall()]
    if 'course' not in cols:
        conn.execute("ALTER TABLE groups_list ADD COLUMN course INTEGER DEFAULT 1")
    if 'is_distance' not in cols:
        conn.execute("ALTER TABLE groups_list ADD COLUMN is_distance INTEGER DEFAULT 0")
    conn.commit()

    # Migration: text fields -> FK columns in lessons
    lesson_cols = [r[1] for r in conn.execute("PRAGMA table_info(lessons)").fetchall()]
    if 'teacher' in lesson_cols and 'teacher_id' not in lesson_cols:
        # Add new FK columns
        conn.execute("ALTER TABLE lessons ADD COLUMN teacher_id INTEGER")
        conn.execute("ALTER TABLE lessons ADD COLUMN subject_id INTEGER")
        conn.execute("ALTER TABLE lessons ADD COLUMN classroom_id INTEGER")

        # Populate from text (exact match)
        conn.execute("""UPDATE lessons SET teacher_id = (
            SELECT t.id FROM teachers t WHERE t.name = lessons.teacher
        ) WHERE teacher != ''""")
        conn.execute("""UPDATE lessons SET subject_id = (
            SELECT s.id FROM subjects s WHERE s.name = lessons.subject
        ) WHERE subject != ''""")
        conn.execute("""UPDATE lessons SET classroom_id = (
            SELECT c.id FROM classrooms_list c WHERE c.name = lessons.classroom
        ) WHERE classroom != ''""")

        # Auto-create missing reference records
        conn.execute("""INSERT INTO teachers (name) SELECT DISTINCT teacher FROM lessons
            WHERE teacher != '' AND teacher_id IS NULL""")
        conn.execute("""UPDATE lessons SET teacher_id = (
            SELECT t.id FROM teachers t WHERE t.name = lessons.teacher
        ) WHERE teacher != '' AND teacher_id IS NULL""")

        conn.execute("""INSERT INTO subjects (name) SELECT DISTINCT subject FROM lessons
            WHERE subject != '' AND subject_id IS NULL""")
        conn.execute("""UPDATE lessons SET subject_id = (
            SELECT s.id FROM subjects s WHERE s.name = lessons.subject
        ) WHERE subject != '' AND subject_id IS NULL""")

        conn.execute("""INSERT INTO classrooms_list (name) SELECT DISTINCT classroom FROM lessons
            WHERE classroom != '' AND classroom_id IS NULL""")
        conn.execute("""UPDATE lessons SET classroom_id = (
            SELECT c.id FROM classrooms_list c WHERE c.name = lessons.classroom
        ) WHERE classroom != '' AND classroom_id IS NULL""")

        # Drop old text columns (first drop any legacy indexes referencing them)
        for idx_name in ['idx_lessons_teacher', 'idx_lessons_subject', 'idx_lessons_classroom']:
            conn.execute(f"DROP INDEX IF EXISTS {idx_name}")
        conn.execute("ALTER TABLE lessons DROP COLUMN teacher")
        conn.execute("ALTER TABLE lessons DROP COLUMN subject")
        conn.execute("ALTER TABLE lessons DROP COLUMN classroom")
        conn.commit()

    # Migration: groups TEXT -> group_id + subgroup
    if 'groups' in lesson_cols and 'group_id' not in lesson_cols:
        conn.execute("ALTER TABLE lessons ADD COLUMN group_id INTEGER")
        conn.execute("ALTER TABLE lessons ADD COLUMN subgroup INTEGER DEFAULT NULL")

        all_groups = conn.execute("SELECT id, groups FROM lessons WHERE groups != ''").fetchall()
        for row in all_groups:
            g = row['groups'] or ''
            m = re.search(r'\((\d+)\)\s*$', g)
            subgroup_val = int(m.group(1)) if m else None
            group_name = g[:m.start()].strip() if m else g.strip()
            if not group_name:
                continue
            existing = conn.execute("SELECT id FROM groups_list WHERE name = ?", (group_name,)).fetchone()
            if existing:
                gid = existing['id']
            else:
                gid = conn.execute("INSERT INTO groups_list (name, course, is_distance) VALUES (?, 1, 0)", (group_name,)).lastrowid
            conn.execute("UPDATE lessons SET group_id = ?, subgroup = ? WHERE id = ?", (gid, subgroup_val, row['id']))

        conn.execute("ALTER TABLE lessons DROP COLUMN groups")
        conn.commit()
    conn.execute("CREATE INDEX IF NOT EXISTS idx_days_group_date ON schedule_days(group_name, date)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_lessons_day_period ON lessons(day_id, period)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_lessons_teacher_id ON lessons(teacher_id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_lessons_subject_id ON lessons(subject_id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_lessons_classroom_id ON lessons(classroom_id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_lessons_group_id ON lessons(group_id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_lessons_deleted_at ON lessons(deleted_at)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_days_deleted_at ON schedule_days(deleted_at)")

    # Migration: excluded_dates table
    existing_tables = [r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()]
    if 'excluded_dates' not in existing_tables:
        conn.execute('''
            CREATE TABLE excluded_dates (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                start_date TEXT NOT NULL,
                end_date TEXT NOT NULL,
                reason TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')

    # Migration: deleted_at columns for soft delete
    lesson_cols = [r[1] for r in conn.execute("PRAGMA table_info(lessons)").fetchall()]
    if 'deleted_at' not in lesson_cols:
        conn.execute("ALTER TABLE lessons ADD COLUMN deleted_at TIMESTAMP NULL")
    day_cols = [r[1] for r in conn.execute("PRAGMA table_info(schedule_days)").fetchall()]
    if 'deleted_at' not in day_cols:
        conn.execute("ALTER TABLE schedule_days ADD COLUMN deleted_at TIMESTAMP NULL")
    conn.commit()

    # Migration: lesson_history table
    if 'lesson_history' not in existing_tables:
        conn.execute('''
            CREATE TABLE lesson_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                lesson_id INTEGER NOT NULL,
                changed_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                old_subject_id INTEGER,
                old_teacher_id INTEGER,
                old_classroom_id INTEGER,
                old_group_id INTEGER,
                old_subgroup INTEGER,
                FOREIGN KEY (lesson_id) REFERENCES lessons(id) ON DELETE CASCADE
            )
        ''')

    # Migration: default_teacher_id on subjects
    subject_cols = [r[1] for r in conn.execute("PRAGMA table_info(subjects)").fetchall()]
    if 'default_teacher_id' not in subject_cols:
        conn.execute("ALTER TABLE subjects ADD COLUMN default_teacher_id INTEGER NULL")

    conn.execute("CREATE INDEX IF NOT EXISTS idx_lesson_history_lesson_id ON lesson_history(lesson_id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_lesson_history_changed_at ON lesson_history(changed_at)")

    conn.commit()
    conn.close()


def is_date_excluded(conn, date_str):
    row = conn.execute(
        "SELECT 1 FROM excluded_dates WHERE start_date <= ? AND end_date >= ? LIMIT 1",
        (date_str, date_str)
    ).fetchone()
    return row is not None


def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get('is_admin'):
            return jsonify({"error": "Unauthorized"}), 401
        return f(*args, **kwargs)
    return decorated


def generic_crud(table, label, order_by='id', extra_fields=None):
    extra_fields = extra_fields or []

    def make_get(label, table, order_by):
        def get_items():
            conn = get_db()
            items = conn.execute(f'SELECT * FROM {table} ORDER BY {order_by}').fetchall()
            conn.close()
            return jsonify([dict(i) for i in items])
        get_items.__name__ = f'get_{label}'
        return get_items

    def make_create(label, table, extra_fields):
        @login_required
        def create_item():
            data = request.get_json()
            name = data.get('name', '').strip()
            if not name:
                return jsonify({"error": "Name required"}), 400
            conn = get_db()
            extra_cols = []
            extra_vals = []
            for f in extra_fields:
                if f in data:
                    extra_cols.append(f)
                    extra_vals.append(data[f])
            cols_sql = ', '.join(['name'] + extra_cols)
            placeholders = ', '.join(['?'] * (1 + len(extra_cols)))
            try:
                cursor = conn.execute(f'INSERT INTO {table} ({cols_sql}) VALUES ({placeholders})', [name] + extra_vals)
                conn.commit()
                item_id = cursor.lastrowid
            except sqlite3.IntegrityError:
                conn.close()
                return jsonify({"error": "Уже существует"}), 409
            conn.close()
            result = {"id": item_id, "name": name}
            for f in extra_fields:
                result[f] = data.get(f)
            return jsonify(result), 201
        create_item.__name__ = f'create_{label}'
        return create_item

    def make_update(label, table, extra_fields):
        @login_required
        def update_item(item_id):
            data = request.get_json()
            name = data.get('name', '').strip()
            if not name:
                return jsonify({"error": "Name required"}), 400
            conn = get_db()
            set_parts = ['name = ?']
            params = [name]
            for f in extra_fields:
                if f in data:
                    set_parts.append(f'{f} = ?')
                    params.append(data[f])
            params.append(item_id)
            conn.execute(f'UPDATE {table} SET {", ".join(set_parts)} WHERE id = ?', params)
            conn.commit()
            conn.close()
            return jsonify({"success": True})
        update_item.__name__ = f'update_{label}'
        return update_item

    def make_delete(label, table):
        @login_required
        def delete_item(item_id):
            conn = get_db()
            conn.execute(f'DELETE FROM {table} WHERE id = ?', (item_id,))
            conn.commit()
            conn.close()
            return jsonify({"success": True})
        delete_item.__name__ = f'delete_{label}'
        return delete_item

    app.add_url_rule(f'/api/{label}', f'get_{label}', make_get(label, table, order_by), methods=['GET'])
    app.add_url_rule(f'/api/{label}', f'create_{label}', make_create(label, table, extra_fields), methods=['POST'])
    app.add_url_rule(f'/api/{label}/<int:item_id>', f'update_{label}', make_update(label, table, extra_fields), methods=['PUT'])
    app.add_url_rule(f'/api/{label}/<int:item_id>', f'delete_{label}', make_delete(label, table), methods=['DELETE'])


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/api/bell-schedule')
def bell_schedule():
    return jsonify(BELL_SCHEDULE)


@app.before_request
def csrf_check():
    if request.method in ('POST', 'PUT', 'DELETE', 'PATCH'):
        if request.path.startswith('/api/'):
            if request.path in ('/api/auth/login', '/api/auth/logout', '/api/auth/check'):
                return
            token = request.headers.get('X-CSRF-Token') or request.headers.get('X-CSRF-TOKEN') or ''
            if token != session.get('csrf_token'):
                return jsonify({"error": "CSRF token missing or invalid"}), 403

@app.route('/api/auth/csrf-token')
def csrf_token():
    if 'csrf_token' not in session:
        session['csrf_token'] = secrets.token_hex(32)
    return jsonify({"csrf_token": session['csrf_token']})


@app.route('/api/auth/login', methods=['POST'])
def login():
    _clean_login_attempts()
    data = request.get_json()
    ip = request.remote_addr
    now = time.time()

    attempts = _login_attempts.get(ip, [])
    attempts = [t for t in attempts if now - t < LOGIN_RATE_WINDOW]
    _login_attempts[ip] = attempts

    if len(attempts) >= LOGIN_RATE_LIMIT:
        return jsonify({"error": "Too many attempts. Try later."}), 429

    if data.get('password') and check_password_hash(ADMIN_PASSWORD_HASH, data['password']):
        session['is_admin'] = True
        session['csrf_token'] = secrets.token_hex(32)
        _login_attempts.pop(ip, None)
        return jsonify({"success": True, "csrf_token": session['csrf_token']})

    attempts.append(now)
    _login_attempts[ip] = attempts
    return jsonify({"error": "Wrong password"}), 401


@app.route('/api/auth/logout', methods=['POST'])
def logout():
    session.pop('is_admin', None)
    return jsonify({"success": True})


@app.route('/api/auth/check')
def check_auth():
    return jsonify({"is_admin": session.get('is_admin', False)})


@app.route('/api/check_date')
def check_date():
    date_str = request.args.get('date', '')
    if not date_str:
        return jsonify({"error": "date parameter required (YYYY-MM-DD)"}), 400
    conn = get_db()
    excluded = is_date_excluded(conn, date_str)
    conn.close()
    return jsonify({"date": date_str, "excluded": excluded})


@app.route('/api/days', methods=['GET'])
def get_days():
    group_name = request.args.get('group', '')
    conn = get_db()
    query = 'SELECT * FROM schedule_days WHERE deleted_at IS NULL'
    params = []
    if group_name:
        query += ' AND group_name = ?'
        params.append(group_name)
    query += ' ORDER BY sort_order'
    days = conn.execute(query, params).fetchall()
    result = []
    for day in days:
        lessons = conn.execute('''
            SELECT l.*, t.name AS teacher_name, s.name AS subject_name, c.name AS classroom_name,
                   CASE WHEN l.subgroup IS NOT NULL THEN gl.name || ' (' || l.subgroup || ')' ELSE gl.name END AS groups
            FROM lessons l
            LEFT JOIN teachers t ON l.teacher_id = t.id
            LEFT JOIN subjects s ON l.subject_id = s.id
            LEFT JOIN classrooms_list c ON l.classroom_id = c.id
            LEFT JOIN groups_list gl ON l.group_id = gl.id
            WHERE l.day_id = ? AND l.deleted_at IS NULL ORDER BY l.period
        ''', (day['id'],)).fetchall()
        result.append({
            "id": day['id'],
            "day_name": day['day_name'],
            "date": day['date'],
            "sort_order": day['sort_order'],
            "group_name": day['group_name'],
            "lessons": [dict(l) for l in lessons]
        })
    conn.close()
    return jsonify(result)


@app.route('/api/teacher-schedule', methods=['GET'])
def get_teacher_schedule():
    teacher_id = request.args.get('teacher_id', '')
    if not teacher_id:
        return jsonify([])
    teacher_id = int(teacher_id)
    conn = get_db()
    lessons = conn.execute('''
        SELECT l.*, sd.day_name, sd.date, sd.group_name,
               t.name AS teacher_name, s.name AS subject_name, c.name AS classroom_name,
               CASE WHEN l.subgroup IS NOT NULL THEN gl.name || ' (' || l.subgroup || ')' ELSE gl.name END AS groups
        FROM lessons l
        JOIN schedule_days sd ON l.day_id = sd.id
        LEFT JOIN teachers t ON l.teacher_id = t.id
        LEFT JOIN subjects s ON l.subject_id = s.id
        LEFT JOIN classrooms_list c ON l.classroom_id = c.id
        LEFT JOIN groups_list gl ON l.group_id = gl.id
        WHERE l.teacher_id = ? AND l.subject_id IS NOT NULL AND l.deleted_at IS NULL
        ORDER BY sd.sort_order, l.period
    ''', (teacher_id,)).fetchall()
    conn.close()

    days_map = {}
    for l in lessons:
        day_key = (l['day_name'], l['date'])
        if day_key not in days_map:
            days_map[day_key] = {
                'day_name': l['day_name'],
                'date': l['date'],
                'lessons': []
            }
        days_map[day_key]['lessons'].append({
            'id': l['id'],
            'period': l['period'],
            'subject': l['subject_name'],
            'teacher': l['teacher_name'],
            'classroom': l['classroom_name'],
            'groups': l['group_name'] or l['groups']
        })

    return jsonify(list(days_map.values()))


@app.route('/api/teachers-with-schedule')
def teachers_with_schedule():
    conn = get_db()
    rows = conn.execute('''
        SELECT DISTINCT t.id, t.name
        FROM lessons l
        JOIN teachers t ON l.teacher_id = t.id
        JOIN schedule_days sd ON l.day_id = sd.id
        WHERE l.subject_id IS NOT NULL AND l.deleted_at IS NULL
        ORDER BY t.name
    ''').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name']} for r in rows])


@app.route('/api/validate', methods=['POST'])
@login_required
def validate_lesson():
    data = request.get_json()
    day_id = data.get('day_id')
    period = data.get('period')
    teacher_id = data.get('teacher_id')
    classroom_id = data.get('classroom_id')
    subject_id = data.get('subject_id')
    group_id = data.get('group_id')
    subgroup = data.get('subgroup')
    exclude_lesson_id = data.get('exclude_lesson_id')

    if not day_id or not period:
        return jsonify({"warnings": [], "blocking": False})

    conn = get_db()
    day = conn.execute('SELECT * FROM schedule_days WHERE id = ? AND deleted_at IS NULL', (day_id,)).fetchall()
    if not day:
        conn.close()
        return jsonify({"warnings": [], "blocking": False})
    day = day[0]

    warnings = []
    blocking = False

    if teacher_id:
        existing = conn.execute('''
            SELECT l.*, sd.day_name, sd.date, sd.group_name, t.name AS teacher_name,
                   CASE WHEN l.subgroup IS NOT NULL THEN gl.name || ' (' || l.subgroup || ')' ELSE gl.name END AS groups
            FROM lessons l
            JOIN schedule_days sd ON l.day_id = sd.id
            LEFT JOIN teachers t ON l.teacher_id = t.id
            LEFT JOIN groups_list gl ON l.group_id = gl.id
            WHERE l.teacher_id = ? AND l.period = ? AND sd.date = ?
            AND l.subject_id IS NOT NULL AND l.deleted_at IS NULL AND sd.deleted_at IS NULL
        ''', (teacher_id, period, day['date'])).fetchall()
        if exclude_lesson_id:
            existing = [e for e in existing if e['id'] != exclude_lesson_id]
        if existing:
            tname = existing[0]['teacher_name'] or ''
            groups = [e['group_name'] or e['groups'] for e in existing]
            warnings.append({
                "type": "teacher_conflict",
                "message": f"Преподаватель {tname} уже занят в {existing[0]['day_name']} на {period}-й паре (группа: {', '.join(groups)})"
            })
            blocking = True

    if classroom_id:
        existing = conn.execute('''
            SELECT l.*, sd.day_name, sd.date, sd.group_name, c.name AS classroom_name,
                   CASE WHEN l.subgroup IS NOT NULL THEN gl.name || ' (' || l.subgroup || ')' ELSE gl.name END AS groups
            FROM lessons l
            JOIN schedule_days sd ON l.day_id = sd.id
            LEFT JOIN classrooms_list c ON l.classroom_id = c.id
            LEFT JOIN groups_list gl ON l.group_id = gl.id
            WHERE l.classroom_id = ? AND l.period = ? AND sd.date = ?
            AND l.subject_id IS NOT NULL AND l.deleted_at IS NULL AND sd.deleted_at IS NULL
        ''', (classroom_id, period, day['date'])).fetchall()
        if exclude_lesson_id:
            existing = [e for e in existing if e['id'] != exclude_lesson_id]
        if existing:
            cname = existing[0]['classroom_name'] or ''
            groups = [e['group_name'] or e['groups'] for e in existing]
            warnings.append({
                "type": "classroom_conflict",
                "message": f"Аудитория {cname} уже занята в {existing[0]['day_name']} на {period}-й паре (группа: {', '.join(groups)})"
            })
            blocking = True

    if group_id is not None and subject_id:
        existing = conn.execute('''
            SELECT l.* FROM lessons l
            WHERE l.day_id = ? AND l.period = ? AND l.subject_id IS NOT NULL AND l.deleted_at IS NULL
        ''', (day_id, period)).fetchall()
        if exclude_lesson_id:
            existing = [e for e in existing if e['id'] != exclude_lesson_id]
        for e in existing:
            e_sub = e['subgroup']
            new_sub = subgroup
            e_gid = e['group_id']
            if e_gid == group_id and (e_sub or '') == (new_sub or ''):
                warnings.append({
                    "type": "duplicate",
                    "message": "Такой же урок уже добавлен в этот день на эту пару"
                })
                blocking = True
                break

    conn.close()
    return jsonify({"warnings": warnings, "blocking": blocking})


@app.route('/api/validate-week', methods=['POST'])
@login_required
def validate_week():
    data = request.get_json()
    group_name = data.get('group_name', '')
    dates = data.get('dates', [])

    if not group_name or not dates:
        return jsonify({"ok": True, "warnings": []})

    conn = get_db()
    warnings = []

    placeholders = ','.join('?' * len(dates))
    existing = conn.execute(
        f'SELECT * FROM schedule_days WHERE group_name = ? AND deleted_at IS NULL AND date IN ({placeholders})',
        [group_name] + dates
    ).fetchall()

    if existing:
        existing_dates = [e['date'] for e in existing]
        warnings.append({
            "type": "week_exists",
            "message": f"У группы {group_name} уже есть расписание на даты: {', '.join(existing_dates)}. Удалите старую неделю перед добавлением новой."
        })

    conn.close()
    return jsonify({"ok": len(warnings) == 0, "warnings": warnings})


@app.route('/api/days', methods=['POST'])
@login_required
def create_day():
    data = request.get_json()
    conn = get_db()
    max_order = conn.execute('SELECT COALESCE(MAX(sort_order), -1) FROM schedule_days').fetchone()[0]
    cursor = conn.execute(
        'INSERT INTO schedule_days (day_name, date, sort_order, group_name) VALUES (?, ?, ?, ?)',
        (data['day_name'], data['date'], max_order + 1, data.get('group_name', ''))
    )
    day_id = cursor.lastrowid
    conn.commit()
    conn.close()
    return jsonify({"id": day_id}), 201


@app.route('/api/days/batch', methods=['POST'])
@login_required
def create_days_batch():
    data = request.get_json()
    days = data.get('days', [])
    if not days:
        return jsonify({"error": "Нет дней"}), 400

    group_name = days[0].get('group_name', '')
    dates = [d['date'] for d in days]

    conn = get_db()

    placeholders = ','.join('?' * len(dates))
    existing = conn.execute(
        f'SELECT date FROM schedule_days WHERE group_name = ? AND deleted_at IS NULL AND date IN ({placeholders})',
        [group_name] + dates
    ).fetchall()

    if existing:
        existing_dates = [e['date'] for e in existing]
        conn.close()
        return jsonify({
            "error": f"У группы {group_name} уже есть расписание на даты: {', '.join(existing_dates)}"
        }), 409

    max_order = conn.execute('SELECT COALESCE(MAX(sort_order), -1) FROM schedule_days').fetchone()[0]
    ids = []
    for i, d in enumerate(days):
        cursor = conn.execute(
            'INSERT INTO schedule_days (day_name, date, sort_order, group_name) VALUES (?, ?, ?, ?)',
            (d['day_name'], d['date'], max_order + 1 + i, d.get('group_name', ''))
        )
        ids.append(cursor.lastrowid)
    conn.commit()
    conn.close()
    return jsonify({"ids": ids}), 201


@app.route('/api/days/<int:day_id>', methods=['PUT'])
@login_required
def update_day(day_id):
    data = request.get_json()
    conn = get_db()
    conn.execute(
        'UPDATE schedule_days SET day_name = ?, date = ? WHERE id = ?',
        (data['day_name'], data['date'], day_id)
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})


@app.route('/api/days/<int:day_id>', methods=['DELETE'])
@login_required
def delete_day(day_id):
    conn = get_db()
    conn.execute("UPDATE schedule_days SET deleted_at = datetime('now') WHERE id = ? AND deleted_at IS NULL", (day_id,))
    conn.execute("UPDATE lessons SET deleted_at = datetime('now') WHERE day_id = ?", (day_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})


def _cleanup_old_deleted():
    conn = get_db()
    conn.execute("DELETE FROM lessons WHERE deleted_at IS NOT NULL AND deleted_at < datetime('now', '-30 days')")
    conn.execute("DELETE FROM schedule_days WHERE deleted_at IS NOT NULL AND deleted_at < datetime('now', '-30 days')")
    conn.commit()
    conn.close()


@app.route('/api/days/<int:day_id>/restore', methods=['POST'])
@login_required
def restore_day(day_id):
    conn = get_db()
    conn.execute("UPDATE schedule_days SET deleted_at = NULL WHERE id = ?", (day_id,))
    conn.execute("UPDATE lessons SET deleted_at = NULL WHERE day_id = ?", (day_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})


def _validate_lesson_data(conn, day_id, period, teacher_id, classroom_id, subject_id=None, exclude_id=None, group_id=None, subgroup=None):
    warnings = []
    blocking = False

    day = conn.execute('SELECT * FROM schedule_days WHERE id = ? AND deleted_at IS NULL', (day_id,)).fetchone()
    if not day:
        return warnings, blocking

    if teacher_id:
        existing = conn.execute('''
            SELECT l.*, sd.day_name, t.name AS teacher_name FROM lessons l
            JOIN schedule_days sd ON l.day_id = sd.id
            LEFT JOIN teachers t ON l.teacher_id = t.id
            WHERE l.teacher_id = ? AND l.period = ? AND sd.date = ? AND l.subject_id IS NOT NULL AND l.deleted_at IS NULL AND sd.deleted_at IS NULL
        ''', (teacher_id, period, day['date'])).fetchall()
        if exclude_id:
            existing = [e for e in existing if e['id'] != exclude_id]
        if existing:
            warnings.append(f"Teacher busy on {existing[0]['day_name']} period {period}")
            blocking = True

    if classroom_id:
        existing = conn.execute('''
            SELECT l.*, sd.day_name, c.name AS classroom_name FROM lessons l
            JOIN schedule_days sd ON l.day_id = sd.id
            LEFT JOIN classrooms_list c ON l.classroom_id = c.id
            WHERE l.classroom_id = ? AND l.period = ? AND sd.date = ? AND l.subject_id IS NOT NULL AND l.deleted_at IS NULL AND sd.deleted_at IS NULL
        ''', (classroom_id, period, day['date'])).fetchall()
        if exclude_id:
            existing = [e for e in existing if e['id'] != exclude_id]
        if existing:
            warnings.append(f"Classroom busy on {existing[0]['day_name']} period {period}")
            blocking = True

    if group_id is not None and subject_id:
        existing = conn.execute('''
            SELECT l.* FROM lessons l
            WHERE l.day_id = ? AND l.period = ? AND l.subject_id IS NOT NULL AND l.deleted_at IS NULL
        ''', (day_id, period)).fetchall()
        if exclude_id:
            existing = [e for e in existing if e['id'] != exclude_id]
        for e in existing:
            e_sub = e['subgroup']
            new_sub = subgroup
            e_gid = e['group_id']
            if e_gid == group_id and (e_sub or '') == (new_sub or ''):
                warnings.append("Duplicate lesson in this cell")
                blocking = True
                break

    return warnings, blocking


@app.route('/api/days/<int:day_id>/lessons', methods=['POST'])
@login_required
def create_lesson(day_id):
    data = request.get_json()
    conn = get_db()
    try:
        conn.execute('BEGIN IMMEDIATE')

        teacher_id = data.get('teacher_id')
        classroom_id = data.get('classroom_id')
        subject_id = data.get('subject_id')
        period = data.get('period')
        if not period or period not in range(1, 9):
            conn.rollback()
            return jsonify({"error": "Invalid period (must be 1-8)"}), 400
        group_id = data.get('group_id')
        subgroup = data.get('subgroup')

        warnings, blocking = _validate_lesson_data(conn, day_id, period, teacher_id, classroom_id, subject_id, group_id=group_id, subgroup=subgroup)
        if blocking:
            conn.rollback()
            return jsonify({"error": "Conflict detected", "warnings": warnings}), 409

        cursor = conn.execute(
            'INSERT INTO lessons (day_id, period, subject_id, teacher_id, classroom_id, group_id, subgroup) VALUES (?, ?, ?, ?, ?, ?, ?)',
            (day_id, period, subject_id, teacher_id, classroom_id, group_id, subgroup)
        )
        lesson_id = cursor.lastrowid
        conn.commit()
        return jsonify({"id": lesson_id}), 201
    finally:
        conn.close()


@app.route('/api/lessons/<int:lesson_id>', methods=['PUT'])
@login_required
def update_lesson(lesson_id):
    data = request.get_json()
    conn = get_db()
    try:
        conn.execute('BEGIN IMMEDIATE')

        existing = conn.execute('SELECT * FROM lessons WHERE id = ? AND deleted_at IS NULL', (lesson_id,)).fetchone()
        if not existing:
            conn.rollback()
            return jsonify({"error": "Lesson not found"}), 404

        teacher_id = data.get('teacher_id')
        classroom_id = data.get('classroom_id')
        subject_id = data.get('subject_id')
        day_id = existing['day_id']
        period = data.get('period', existing['period'])
        group_id = data.get('group_id')
        subgroup = data.get('subgroup')

        warnings, blocking = _validate_lesson_data(conn, day_id, period, teacher_id, classroom_id, subject_id, exclude_id=lesson_id, group_id=group_id, subgroup=subgroup)
        if blocking:
            conn.rollback()
            return jsonify({"error": "Conflict detected", "warnings": warnings}), 409

        conn.execute(
            'INSERT INTO lesson_history (lesson_id, old_subject_id, old_teacher_id, old_classroom_id, old_group_id, old_subgroup) VALUES (?, ?, ?, ?, ?, ?)',
            (lesson_id, existing['subject_id'], existing['teacher_id'], existing['classroom_id'], existing.get('group_id'), existing.get('subgroup'))
        )

        conn.execute(
            'UPDATE lessons SET subject_id = ?, teacher_id = ?, classroom_id = ?, group_id = ?, subgroup = ? WHERE id = ?',
            (subject_id, teacher_id, classroom_id, group_id, subgroup, lesson_id)
        )
        conn.commit()
        return jsonify({"success": True})
    finally:
        conn.close()


@app.route('/api/lessons/<int:lesson_id>', methods=['DELETE'])
@login_required
def delete_lesson(lesson_id):
    conn = get_db()
    conn.execute("UPDATE lessons SET deleted_at = datetime('now') WHERE id = ? AND deleted_at IS NULL", (lesson_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})


@app.route('/api/lessons/<int:lesson_id>/restore', methods=['POST'])
@login_required
def restore_lesson(lesson_id):
    conn = get_db()
    conn.execute("UPDATE lessons SET deleted_at = NULL WHERE id = ?", (lesson_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})


@app.route('/api/trash')
@login_required
def get_trash():
    conn = get_db()
    deleted_days = conn.execute('''
        SELECT d.id, d.day_name, d.date, d.group_name, d.deleted_at
        FROM schedule_days d
        WHERE d.deleted_at IS NOT NULL AND d.deleted_at >= datetime('now', '-30 days')
        ORDER BY d.deleted_at DESC
    ''').fetchall()
    deleted_lessons = conn.execute('''
        SELECT l.id, l.day_id, l.period, l.deleted_at,
               t.name AS teacher_name, s.name AS subject_name, c.name AS classroom_name,
               gl.name AS group_name, l.subgroup,
               sd.day_name, sd.date
        FROM lessons l
        JOIN schedule_days sd ON l.day_id = sd.id
        LEFT JOIN teachers t ON l.teacher_id = t.id
        LEFT JOIN subjects s ON l.subject_id = s.id
        LEFT JOIN classrooms_list c ON l.classroom_id = c.id
        LEFT JOIN groups_list gl ON l.group_id = gl.id
        WHERE l.deleted_at IS NOT NULL AND l.deleted_at >= datetime('now', '-30 days')
        ORDER BY l.deleted_at DESC
    ''').fetchall()
    conn.close()
    return jsonify({
        "days": [dict(d) for d in deleted_days],
        "lessons": [dict(l) for l in deleted_lessons]
    })


@app.route('/api/classrooms-with-schedule')
def classrooms_with_schedule():
    conn = get_db()
    rows = conn.execute('''
        SELECT DISTINCT c.id, c.name
        FROM lessons l
        JOIN classrooms_list c ON l.classroom_id = c.id
        JOIN schedule_days sd ON l.day_id = sd.id
        WHERE l.subject_id IS NOT NULL AND l.deleted_at IS NULL
        ORDER BY CAST(c.name AS INTEGER), c.name
    ''').fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name']} for r in rows])


def build_week_dates(monday_str):
    monday = datetime.strptime(monday_str, '%Y-%m-%d').date()
    names = ['Понедельник', 'Вторник', 'Среда', 'Четверг', 'Пятница', 'Суббота']
    return [{'day_name': names[i], 'date': (monday + timedelta(days=i)).isoformat()} for i in range(6)]


def _fetch_group_raw(group_name, monday):
    week_dates = build_week_dates(monday)
    date_list = [d['date'] for d in week_dates]
    conn = get_db()
    placeholders = ','.join('?' * len(date_list))
    days = conn.execute(
        f'SELECT * FROM schedule_days WHERE group_name = ? AND deleted_at IS NULL AND date IN ({placeholders}) ORDER BY date',
        [group_name] + date_list
    ).fetchall()
    day_id_map = {day['date']: day['id'] for day in days}
    for wd in week_dates:
        wd['day_id'] = day_id_map.get(wd['date'])
    lessons_by_date = {}
    for day in days:
        lessons = conn.execute('''
            SELECT l.*, t.name AS teacher_name, s.name AS subject_name, c.name AS classroom_name,
                   CASE WHEN l.subgroup IS NOT NULL THEN gl.name || ' (' || l.subgroup || ')' ELSE gl.name END AS groups
            FROM lessons l
            LEFT JOIN teachers t ON l.teacher_id = t.id
            LEFT JOIN subjects s ON l.subject_id = s.id
            LEFT JOIN classrooms_list c ON l.classroom_id = c.id
            LEFT JOIN groups_list gl ON l.group_id = gl.id
            WHERE l.day_id = ? AND l.deleted_at IS NULL ORDER BY l.period
        ''', (day['id'],)).fetchall()
        lessons_by_date[day['date']] = {}
        for l in lessons:
            p = l['period']
            if p not in lessons_by_date[day['date']]:
                lessons_by_date[day['date']][p] = []
            lessons_by_date[day['date']][p].append(dict(l))
    conn.close()
    return week_dates, lessons_by_date


def _fetch_entity_raw(entity_table, entity_id, entity_col, monday):
    week_dates = build_week_dates(monday)
    date_list = [d['date'] for d in week_dates]
    conn = get_db()
    placeholders = ','.join('?' * len(date_list))
    rows = conn.execute(f'''
        SELECT l.*, sd.date, sd.group_name, t.name AS teacher_name,
               s.name AS subject_name, c.name AS classroom_name,
               CASE WHEN l.subgroup IS NOT NULL THEN gl.name || ' (' || l.subgroup || ')' ELSE gl.name END AS groups
        FROM lessons l
        JOIN schedule_days sd ON l.day_id = sd.id
        LEFT JOIN teachers t ON l.teacher_id = t.id
        LEFT JOIN subjects s ON l.subject_id = s.id
        LEFT JOIN classrooms_list c ON l.classroom_id = c.id
        LEFT JOIN groups_list gl ON l.group_id = gl.id
        WHERE l.{entity_col} = ? AND l.subject_id IS NOT NULL AND l.deleted_at IS NULL AND sd.deleted_at IS NULL AND sd.date IN ({placeholders})
        ORDER BY sd.date, l.period
    ''', [entity_id] + date_list).fetchall()
    conn.close()
    lessons_by_date_period = {}
    day_id_map = {}
    for r in rows:
        d = r['date']
        p = r['period']
        day_id_map[d] = r['day_id']
        if d not in lessons_by_date_period:
            lessons_by_date_period[d] = {}
        lessons_by_date_period[d][p] = {
            'id': r['id'], 'subject': r['subject_name'], 'teacher': r['teacher_name'],
            'classroom': r['classroom_name'],
            'groups': r['group_name'] or r['groups'],
            'day_id': r['day_id'],
            'group_id': r['group_id'], 'subgroup': r['subgroup']
        }
    for wd in week_dates:
        wd['day_id'] = day_id_map.get(wd['date'])
    return week_dates, lessons_by_date_period


@app.route('/api/schedule-table')
def schedule_table():
    group_name = request.args.get('group', '')
    monday = request.args.get('monday', '')
    if not group_name or not monday:
        return jsonify({'days': [], 'periods': []})

    week_dates, lessons_by_date = _fetch_group_raw(group_name, monday)

    periods = []
    for b in BELL_SCHEDULE:
        p = b['period']
        lessons = {}
        for wd in week_dates:
            d = wd['date']
            if d in lessons_by_date and p in lessons_by_date[d]:
                raw_list = lessons_by_date[d][p]
                lessons[d] = [{
                    'id': l['id'], 'subject': l['subject_name'], 'teacher': l['teacher_name'],
                    'classroom': l['classroom_name'], 'groups': l['groups'],
                    'day_id': l['day_id'],
                    'subject_id': l['subject_id'], 'teacher_id': l['teacher_id'],
                    'classroom_id': l['classroom_id'],
                    'group_id': l['group_id'], 'subgroup': l['subgroup']
                } for l in raw_list]
            else:
                lessons[d] = []
        periods.append({**b, 'lessons': lessons})

    return jsonify({'days': week_dates, 'periods': periods, 'week_exists': bool(lessons_by_date)})


def _build_schedule_periods(week_dates, lessons_by_date_period):
    periods = []
    for b in BELL_SCHEDULE:
        p = b['period']
        lessons = {}
        for wd in week_dates:
            d = wd['date']
            lesson = lessons_by_date_period.get(d, {}).get(p)
            if lesson:
                lessons[d] = [{
                    'id': lesson['id'], 'subject': lesson['subject'], 'teacher': lesson['teacher'],
                    'classroom': lesson['classroom'], 'groups': lesson['groups'],
                    'day_id': lesson['day_id'],
                    'group_id': lesson.get('group_id'), 'subgroup': lesson.get('subgroup')
                }]
            else:
                lessons[d] = []
        periods.append({**b, 'lessons': lessons})
    return periods


@app.route('/api/teacher-schedule-table')
def teacher_schedule_table():
    teacher_id = request.args.get('teacher_id', '')
    monday = request.args.get('monday', '')
    if not teacher_id or not monday:
        return jsonify({'days': [], 'periods': []})
    teacher_id = int(teacher_id)

    week_dates, lessons_by_date_period = _fetch_entity_raw('teachers', teacher_id, 'teacher_id', monday)

    return jsonify({'days': week_dates, 'periods': _build_schedule_periods(week_dates, lessons_by_date_period)})


@app.route('/api/classroom-schedule-table')
def classroom_schedule_table():
    classroom_id = request.args.get('classroom_id', '')
    monday = request.args.get('monday', '')
    if not classroom_id or not monday:
        return jsonify({'days': [], 'periods': []})
    classroom_id = int(classroom_id)

    week_dates, lessons_by_date_period = _fetch_entity_raw('classrooms_list', classroom_id, 'classroom_id', monday)

    return jsonify({'days': week_dates, 'periods': _build_schedule_periods(week_dates, lessons_by_date_period)})


generic_crud('subjects', 'subjects', 'name', extra_fields=['default_teacher_id'])
generic_crud('teachers', 'teachers', 'name')
generic_crud('classrooms_list', 'classrooms', 'CAST(name AS INTEGER), name')


@app.route('/api/groups', methods=['GET'])
def get_groups():
    conn = get_db()
    items = conn.execute('SELECT * FROM groups_list ORDER BY course, name').fetchall()
    conn.close()
    return jsonify([dict(i) for i in items])


@app.route('/api/groups', methods=['POST'])
@login_required
def create_group():
    data = request.get_json()
    name = data.get('name', '').strip()
    if not name:
        return jsonify({"error": "Name required"}), 400
    course = data.get('course', 1)
    is_distance = 1 if data.get('is_distance') else 0
    conn = get_db()
    existing = conn.execute("SELECT id FROM groups_list WHERE LOWER(name) = LOWER(?)", (name,)).fetchone()
    if existing:
        conn.close()
        return jsonify({"error": "Группа с таким названием уже существует"}), 400
    try:
        cursor = conn.execute(
            'INSERT INTO groups_list (name, course, is_distance) VALUES (?, ?, ?)',
            (name, course, is_distance)
        )
        conn.commit()
        item_id = cursor.lastrowid
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": "Уже существует"}), 409
    conn.close()
    return jsonify({"id": item_id, "name": name, "course": course, "is_distance": is_distance}), 201


@app.route('/api/groups/<int:item_id>', methods=['PUT'])
@login_required
def update_group(item_id):
    data = request.get_json()
    name = data.get('name', '').strip()
    if not name:
        return jsonify({"error": "Name required"}), 400
    course = data.get('course', 1)
    is_distance = 1 if data.get('is_distance') else 0
    conn = get_db()
    existing = conn.execute(
        "SELECT id FROM groups_list WHERE LOWER(name) = LOWER(?) AND id != ?",
        (name, item_id)
    ).fetchone()
    if existing:
        conn.close()
        return jsonify({"error": "Группа с таким названием уже существует"}), 400
    conn.execute(
        'UPDATE groups_list SET name = ?, course = ?, is_distance = ? WHERE id = ?',
        (name, course, is_distance, item_id)
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})


@app.route('/api/groups/<int:item_id>', methods=['DELETE'])
@login_required
def delete_group(item_id):
    conn = get_db()
    conn.execute('DELETE FROM groups_list WHERE id = ?', (item_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})


@app.route('/api/groups/advance-courses', methods=['POST'])
@login_required
def advance_courses():
    conn = get_db()
    deleted = conn.execute('DELETE FROM groups_list WHERE course = 4').rowcount
    conn.execute('UPDATE groups_list SET course = course + 1 WHERE course < 4')
    conn.commit()
    conn.close()
    return jsonify({"success": True, "deleted": deleted})


def _safe_filename(s):
    return re.sub(r'[\\/:*?"<>|]+', '_', s).strip('_')


def _get_export_data(schedule_type, name, monday):
    if schedule_type == 'group':
        week_dates, lessons_by_date = _fetch_group_raw(name, monday)
        periods = []
        for b in BELL_SCHEDULE:
            p = b['period']
            lessons = {}
            for wd in week_dates:
                d = wd['date']
                if d in lessons_by_date and p in lessons_by_date[d]:
                    raw_list = lessons_by_date[d][p]
                    lessons[d] = '\n---\n'.join(
                        f"{l['subject_name'] or ''}\n{l['teacher_name'] or ''}\n{l['classroom_name'] or ''}"
                        for l in raw_list
                    )
                else:
                    lessons[d] = ''
            periods.append({**b, 'lessons': lessons})
        return week_dates, periods

    elif schedule_type == 'teacher':
        week_dates, ldp = _fetch_entity_raw('teachers', name, 'teacher_id', monday)
        periods = []
        for b in BELL_SCHEDULE:
            p = b['period']
            lessons = {}
            for wd in week_dates:
                l = ldp.get(wd['date'], {}).get(p)
                lessons[wd['date']] = f"{l['subject'] or ''}\n{l['groups'] or ''}\n{l['classroom'] or ''}" if l else ''
            periods.append({**b, 'lessons': lessons})
        return week_dates, periods

    elif schedule_type == 'classroom':
        week_dates, ldp = _fetch_entity_raw('classrooms_list', name, 'classroom_id', monday)
        periods = []
        for b in BELL_SCHEDULE:
            p = b['period']
            lessons = {}
            for wd in week_dates:
                l = ldp.get(wd['date'], {}).get(p)
                lessons[wd['date']] = f"{l['subject'] or ''}\n{l['teacher'] or ''}\n{l['groups'] or ''}" if l else ''
            periods.append({**b, 'lessons': lessons})
        return week_dates, periods

    return build_week_dates(monday), []


def _export_excel(week_dates, periods, title):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = 'Расписание'

    thin = Side(style='thin', color='999999')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(name='Arial', bold=True, size=11)
    header_fill = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')
    cell_font = Font(name='Arial', size=10)
    wrap = Alignment(wrap_text=True, vertical='top')

    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(week_dates) + 1)
    title_cell = ws.cell(row=row, column=1, value=title)
    title_cell.font = Font(name='Arial', bold=True, size=14)
    row += 1

    ws.cell(row=row, column=1, value='Пара').font = header_font
    ws.cell(row=row, column=1).fill = header_fill
    ws.cell(row=row, column=1).border = border
    ws.cell(row=row, column=1).alignment = Alignment(horizontal='center')

    for i, wd in enumerate(week_dates):
        d = datetime.strptime(wd['date'], '%Y-%m-%d')
        header_text = f"{wd['day_name']}\n{d.day}.{d.month}"
        cell = ws.cell(row=row, column=i + 2, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
    row += 1

    last_shift = 0
    for p in periods:
        if p['shift'] != last_shift:
            last_shift = p['shift']
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(week_dates) + 1)
            shift_cell = ws.cell(row=row, column=1, value=SHIFT_LABELS[p['shift']])
            shift_cell.font = Font(name='Arial', bold=True, size=10, color='1A5FB4')
            shift_cell.fill = PatternFill(start_color='E0ECFF', end_color='E0ECFF', fill_type='solid')
            shift_cell.alignment = Alignment(horizontal='center')
            for c in range(1, len(week_dates) + 2):
                ws.cell(row=row, column=c).border = border
            row += 1

        time_cell = ws.cell(row=row, column=1, value=f"{p['period']} пара\n{p['start']}–{p['end']}")
        time_cell.font = cell_font
        time_cell.border = border
        time_cell.alignment = Alignment(horizontal='center', vertical='top')

        for i, wd in enumerate(week_dates):
            cell = ws.cell(row=row, column=i + 2, value=p['lessons'].get(wd['date'], ''))
            cell.font = cell_font
            cell.border = border
            cell.alignment = wrap
        row += 1

    ws.column_dimensions['A'].width = 16
    for i in range(len(week_dates)):
        letter = get_column_letter(i + 2)
        ws.column_dimensions[letter].width = 28

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _export_word(week_dates, periods, title):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii', 'Arial')
    style.element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi', 'Arial')

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(16)

    num_cols = len(week_dates) + 1
    table = doc.add_table(rows=1, cols=num_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0]
    hdr.cells[0].text = 'Пара'
    for p in hdr.cells[0].paragraphs:
        for r in p.runs:
            r.bold = True
    for i, wd in enumerate(week_dates):
        d = datetime.strptime(wd['date'], '%Y-%m-%d')
        cell = hdr.cells[i + 1]
        cell.text = f"{wd['day_name']}\n{d.day}.{d.month}"
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True

    last_shift = 0
    for p in periods:
        if p['shift'] != last_shift:
            last_shift = p['shift']
            shift_row = table.add_row()
            shift_row.cells[0].merge(shift_row.cells[len(week_dates)])
            shift_row.cells[0].text = SHIFT_LABELS[p['shift']]
            for para in shift_row.cells[0].paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x1A, 0x5F, 0xB4)

        row_cells = table.add_row().cells
        row_cells[0].text = f"{p['period']} пара\n{p['start']}–{p['end']}"
        for i, wd in enumerate(week_dates):
            row_cells[i + 1].text = p['lessons'].get(wd['date'], '')

    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs:
                    run.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _export_pdf(week_dates, periods, title):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_path = os.environ.get('PDF_FONT_PATH', '')
    if not font_path:
        candidates = [
            r'C:\Windows\Fonts\tahoma.ttf',
            r'C:\Windows\Fonts\arial.ttf',
            r'C:\Windows\Fonts\segoeui.ttf',
            '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
            '/usr/share/fonts/TTF/DejaVuSans.ttf',
        ]
        for cp in candidates:
            if os.path.exists(cp):
                font_path = cp
                break
    if not font_path or not os.path.exists(font_path):
        font_path = None

    font_bold_path = os.environ.get('PDF_FONT_BOLD_PATH', '')
    if not font_bold_path and font_path:
        bname = font_path.replace('.ttf', 'bd.ttf')
        if os.path.exists(bname):
            font_bold_path = bname

    if font_path:
        pdfmetrics.registerFont(TTFont('Tahoma', font_path))
        pdfmetrics.registerFont(TTFont('Tahoma-Bold', font_bold_path or font_path))
    font_name = 'Tahoma' if font_path else 'Helvetica'
    font_bold_name = 'Tahoma-Bold' if font_path else 'Helvetica-Bold'

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)

    style_title = ParagraphStyle('Title', fontName=font_bold_name, fontSize=14, alignment=1, spaceAfter=12)
    style_cell = ParagraphStyle('Cell', fontName=font_name, fontSize=8, leading=10)
    style_header = ParagraphStyle('Header', fontName=font_bold_name, fontSize=9, leading=11, alignment=1)

    elements = []
    elements.append(Paragraph(title, style_title))
    elements.append(Spacer(1, 0.3*cm))

    header_dates = [(wd['day_name'], datetime.strptime(wd['date'], '%Y-%m-%d')) for wd in week_dates]
    data = [['Пара'] + [f"{name}\n{d.day}.{d.month}" for name, d in header_dates]]

    last_shift = 0
    for p in periods:
        if p['shift'] != last_shift:
            last_shift = p['shift']
            shift_row_data = [Paragraph(f"<b>{SHIFT_LABELS[p['shift']]}</b>", style_header)]
            shift_row_data += [''] * len(week_dates)
            data.append(shift_row_data)
        row = [Paragraph(f"<b>{p['period']} пара</b><br/>{p['start']}–{p['end']}", style_cell)]
        for wd in week_dates:
            text = p['lessons'].get(wd['date'], '')
            row.append(Paragraph(text.replace('\n', '<br/>'), style_cell) if text else '')
        data.append(row)

    col_widths = [2.5*cm] + [max(2.2*cm, (landscape(A4)[0] - 4*cm - 2.5*cm) / len(week_dates))] * len(week_dates)
    t = Table(data, colWidths=col_widths, repeatRows=1)

    style_cmds = [
        ('FONTNAME', (0, 0), (-1, 0), font_bold_name),
        ('FONTNAME', (0, 1), (-1, -1), font_name),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#D9E2F3')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#999999')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]

    row_idx = 1
    last_shift = 0
    for p in periods:
        if p['shift'] != last_shift:
            last_shift = p['shift']
            style_cmds.append(('SPAN', (0, row_idx), (-1, row_idx)))
            style_cmds.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor('#E0ECFF')))
            style_cmds.append(('ALIGN', (0, row_idx), (-1, row_idx), 'CENTER'))
            row_idx += 1
        row_idx += 1

    t.setStyle(TableStyle(style_cmds))
    elements.append(t)
    doc.build(elements)
    buf.seek(0)
    return buf


@app.route('/api/export')
def export_schedule():
    schedule_type = request.args.get('type', '')
    name = request.args.get('name', '')
    monday = request.args.get('monday', '')
    fmt = request.args.get('format', '')

    if not all([schedule_type, name, monday, fmt]):
        return jsonify({"error": "Missing parameters"}), 400
    if fmt not in ('xlsx', 'docx', 'pdf'):
        return jsonify({"error": "Invalid format"}), 400

    week_dates, periods = _get_export_data(schedule_type, name, monday)

    if not periods:
        return jsonify({"error": "No schedule data"}), 404

    sat = datetime.strptime(week_dates[-1]['date'], '%Y-%m-%d')
    mon = datetime.strptime(week_dates[0]['date'], '%Y-%m-%d')
    date_range = f"{mon.day:02d}.{mon.month:02d}–{sat.day:02d}.{sat.month:02d}"

    type_labels = {'group': 'Группа', 'teacher': 'Преподаватель', 'classroom': 'Аудитория'}
    title = f"Расписание — {type_labels.get(schedule_type, '')} {name}, неделя {date_range}"
    safe_name = _safe_filename(f"raspisanie_{name}_{date_range}")

    if fmt == 'xlsx':
        buf = _export_excel(week_dates, periods, title)
        return send_file(buf, as_attachment=True, download_name=f"{safe_name}.xlsx",
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    elif fmt == 'docx':
        buf = _export_word(week_dates, periods, title)
        return send_file(buf, as_attachment=True, download_name=f"{safe_name}.docx",
                         mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    elif fmt == 'pdf':
        buf = _export_pdf(week_dates, periods, title)
        return send_file(buf, as_attachment=True, download_name=f"{safe_name}.pdf",
                         mimetype='application/pdf')


@app.route('/api/availability')
def availability():
    monday = request.args.get('monday', '')
    if not monday:
        return jsonify({"error": "monday required"}), 400

    week_dates = build_week_dates(monday)
    date_list = [wd['date'] for wd in week_dates]
    placeholders = ','.join('?' * len(date_list))

    conn = get_db()
    try:
        lessons = conn.execute(f'''
            SELECT l.period, sd.date, sd.day_name,
                   t.name AS teacher_name, l.teacher_id,
                   c.name AS classroom_name, l.classroom_id,
                   g.name AS group_name, l.group_id,
                   s.name AS subject_name, l.subject_id,
                   s.lesson_type
            FROM lessons l
            JOIN schedule_days sd ON l.day_id = sd.id
            LEFT JOIN teachers t ON l.teacher_id = t.id
            LEFT JOIN classrooms_list c ON l.classroom_id = c.id
            LEFT JOIN groups_list g ON l.group_id = g.id
            LEFT JOIN subjects s ON l.subject_id = s.id
            WHERE sd.date IN ({placeholders})
              AND l.subject_id IS NOT NULL
              AND l.deleted_at IS NULL AND sd.deleted_at IS NULL
        ''', date_list).fetchall()

        all_teachers = conn.execute('SELECT id, name FROM teachers ORDER BY name').fetchall()
        all_classrooms = conn.execute('SELECT id, name FROM classrooms_list ORDER BY CAST(name AS INTEGER), name').fetchall()

        teacher_busy = {}
        classroom_busy = {}
        for l in lessons:
            key = (l['date'], l['period'])
            t_id = l['teacher_id']
            c_id = l['classroom_id']
            info = {'group': l['group_name'], 'subject': l['subject_name'], 'teacher': l['teacher_name'], 'classroom': l['classroom_name'], 'lesson_type': l['lesson_type'] or ''}
            if t_id:
                teacher_busy.setdefault(t_id, {})[key] = info
            if c_id:
                classroom_busy.setdefault(c_id, {})[key] = info

        periods = [{'period': b['period'], 'start': b['start'], 'end': b['end']} for b in BELL_SCHEDULE]

        return jsonify({
            'days': week_dates,
            'periods': periods,
            'teachers': [{'id': t['id'], 'name': t['name']} for t in all_teachers],
            'classrooms': [{'id': c['id'], 'name': c['name']} for c in all_classrooms],
            'teacher_busy': {str(k): {f"{dk[0]}_{dk[1]}": v for dk, v in dv.items()} for k, dv in teacher_busy.items()},
            'classroom_busy': {str(k): {f"{dk[0]}_{dk[1]}": v for dk, v in dv.items()} for k, dv in classroom_busy.items()},
        })
    finally:
        conn.close()


BACKUPS_DIR = os.path.join(os.path.dirname(__file__), 'backups')


def _daily_backup():
    os.makedirs(BACKUPS_DIR, exist_ok=True)
    today = datetime.now().strftime('%Y-%m-%d')
    existing = [f for f in os.listdir(BACKUPS_DIR)
                if f.startswith(f'schedule_backup_{today}') and f.endswith('.db')]
    if not existing:
        ts = datetime.now().strftime('%Y-%m-%d_%H-%M')
        shutil.copy2(DB_PATH, os.path.join(BACKUPS_DIR, f'schedule_backup_{ts}.db'))
    cutoff = (datetime.now() - timedelta(days=30)).strftime('%Y-%m-%d')
    for f in os.listdir(BACKUPS_DIR):
        if f.startswith('schedule_backup_') and f.endswith('.db'):
            date_part = f.replace('schedule_backup_', '').split('_')[0]
            if date_part < cutoff:
                os.remove(os.path.join(BACKUPS_DIR, f))


if __name__ == '__main__':
    _daily_backup()
    init_db()
    _cleanup_old_deleted()
    debug = os.environ.get('FLASK_DEBUG', '').lower() in ('1', 'true', 'yes')
    host = os.environ.get('FLASK_HOST', '127.0.0.1')
    app.run(debug=debug, host=host, port=5000)
