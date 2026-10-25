"""Unit-тесты модуля «Универсальный импорт справочников из Excel (.xlsx)»."""
import io
import pytest
from openpyxl import Workbook, load_workbook


def _make_xlsx(header, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _post_import(client, entity_type, buf, filename='import.xlsx', csrf='test-csrf'):
    return client.post('/api/admin/import/excel',
                       data={'entity_type': entity_type, 'file': (buf, filename)},
                       headers={'X-CSRF-Token': csrf},
                       content_type='multipart/form-data')


class TestTemplateDownload:
    ENTITIES = ['teachers', 'classrooms', 'groups', 'subjects']

    def test_templates_download_all(self, client, admin_csrf):
        for e in self.ENTITIES:
            resp = client.get(f'/api/admin/import/template/{e}', headers={'X-CSRF-Token': admin_csrf})
            assert resp.status_code == 200
            assert 'spreadsheetml' in resp.content_type
            assert f'template_{e}.xlsx' in resp.headers.get('Content-Disposition', '')

    def test_template_has_header_and_examples(self, client, admin_csrf):
        resp = client.get('/api/admin/import/template/teachers', headers={'X-CSRF-Token': admin_csrf})
        wb = load_workbook(io.BytesIO(resp.data))
        ws = wb.active
        assert ws.max_row >= 2
        header = [c.value for c in ws[1]]
        assert 'ФИО преподавателя' in header

    def test_templates_are_canonical_no_synonyms(self, client, admin_csrf):
        expected = {
            'teachers': ['ФИО преподавателя', 'Телефон'],
            'classrooms': ['Номер аудитории'],
            'subjects': ['Название предмета'],
            'groups': ['Название группы', 'Курс (1-4)', 'Код специальности'],
        }
        for e, want in expected.items():
            resp = client.get(f'/api/admin/import/template/{e}', headers={'X-CSRF-Token': admin_csrf})
            ws = load_workbook(io.BytesIO(resp.data)).active
            header = [c.value for c in ws[1]]
            assert header == want, f'{e}: {header} != {want}'

    def test_template_unknown_entity(self, client, admin_csrf):
        resp = client.get('/api/admin/import/template/unknown', headers={'X-CSRF-Token': admin_csrf})
        assert resp.status_code == 400


class TestImportTeachers:
    def test_import_with_canonical_template_header(self, client, admin_csrf):
        buf = _make_xlsx(['ФИО преподавателя', 'Телефон'], [
            ['Канонический К.К.', '+7-111-222-33-44'],
            ['', ''],
        ])
        resp = _post_import(client, 'teachers', buf)
        assert resp.status_code == 200
        data = resp.get_json()
        assert data['created'] == 1
        items = client.get('/api/teachers').get_json()
        assert any(t['name'] == 'Канонический К.К.' and t['phone'] == '+7-111-222-33-44' for t in items)

    def test_import_creates_and_skips_duplicates(self, client, admin_csrf):
        buf = _make_xlsx(['ФИО', 'Телефон'], [
            ['Иванов И.И.', '+7-900-000-00-00'],
            ['Иванов И.И.', '+7-900-000-00-00'],  # дубль в файле
            ['Петров П.П.', ''],
            ['', ''],  # пустая строка
        ])
        resp = _post_import(client, 'teachers', buf)
        assert resp.status_code == 200
        data = resp.get_json()
        assert data['created'] == 2
        assert data['skipped'] == 1
        items = client.get('/api/teachers').get_json()
        assert len(items) == 2
        assert any(t['phone'] == '+7-900-000-00-00' for t in items)

    def test_import_skips_existing_in_db(self, client, seed_refs, admin_csrf):
        # seed_refs создаёт Иванов И.И.
        buf = _make_xlsx(['ФИО'], [['Иванов И.И.'], ['Новый Н.Н.']])
        resp = _post_import(client, 'teachers', buf)
        assert resp.status_code == 200
        data = resp.get_json()
        assert data['created'] == 1
        assert data['skipped'] == 1


class TestImportGroupsWithSpecialty:
    def test_import_resolves_specialty_by_code(self, client, admin_csrf, db):
        db.seed_specialties()
        buf = _make_xlsx(['Группа', 'Курс', 'Специальность'], [
            ['ЮР-21/26', 1, '40.02.04'],
            ['ЭК-11/26', 2, '38.02.01'],
        ])
        resp = _post_import(client, 'groups', buf)
        assert resp.status_code == 200
        data = resp.get_json()
        assert data['created'] == 2
        assert data['errors'] == []
        groups = client.get('/api/groups').get_json()
        yur = next(g for g in groups if g['name'] == 'ЮР-21/26')
        assert yur['specialty_id'] is not None
        assert yur['specialty_code'] == '40.02.04'
        ek = next(g for g in groups if g['name'] == 'ЭК-11/26')
        assert ek['specialty_code'] == '38.02.01'

    def test_import_group_unknown_specialty_keeps_group(self, client, admin_csrf):
        buf = _make_xlsx(['Группа', 'Курс', 'Специальность'], [['ГР-1/26', 1, '99.99.99']])
        resp = _post_import(client, 'groups', buf)
        assert resp.status_code == 200
        data = resp.get_json()
        assert data['created'] == 1
        assert any('99.99.99' in e for e in data['errors'])
        groups = client.get('/api/groups').get_json()
        assert any(g['name'] == 'ГР-1/26' for g in groups)


class TestImportErrors:
    def test_unknown_entity_type(self, client, admin_csrf):
        resp = _post_import(client, 'nope', _make_xlsx(['X'], [['y']]))
        assert resp.status_code == 400

    def test_wrong_file_extension(self, client, admin_csrf):
        resp = _post_import(client, 'teachers', _make_xlsx(['ФИО'], [['A']]), filename='import.csv')
        assert resp.status_code == 400

    def test_wrong_headers(self, client, admin_csrf):
        buf = _make_xlsx(['Неизвестное поле', 'Что-то'], [['A', 'B']])
        resp = _post_import(client, 'teachers', buf)
        assert resp.status_code == 400
        assert 'названием' in resp.get_json()['error']

    def test_empty_file(self, client, admin_csrf):
        buf = io.BytesIO()
        wb = Workbook()
        wb.save(buf)
        buf.seek(0)
        resp = _post_import(client, 'teachers', buf, filename='empty.xlsx')
        assert resp.status_code == 400

    def test_not_an_xlsx_file(self, client, admin_csrf):
        buf = io.BytesIO(b'not an xlsx content')
        resp = _post_import(client, 'teachers', buf, filename='bad.xlsx')
        assert resp.status_code == 400


class TestImportSubjectsClassrooms:
    def test_import_subjects(self, client, admin_csrf):
        buf = _make_xlsx(['Предмет'], [['Математика'], ['Физика']])
        resp = _post_import(client, 'subjects', buf)
        assert resp.status_code == 200
        data = resp.get_json()
        assert data['created'] == 2
        names = [s['name'] for s in client.get('/api/subjects').get_json()]
        assert 'Математика' in names

    def test_import_classrooms(self, client, admin_csrf):
        buf = _make_xlsx(['Аудитория'], [['101'], ['202']])
        resp = _post_import(client, 'classrooms', buf)
        assert resp.status_code == 200
        assert resp.get_json()['created'] == 2
        names = [c['name'] for c in client.get('/api/classrooms').get_json()]
        assert '101' in names and '202' in names
